from __future__ import annotations

import argparse
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .utils import clean_text, ensure_directory, normalize_identifier, normalize_isin, normalize_text, strip_legal_form, to_float


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_HELPER_WORKBOOK = PROJECT_ROOT / "data_processed" / "cbonds_manual_helper.xlsx"
DEFAULT_CAPTURE_GLOB = "cbonds_capture_output_bucket_*.xlsx"
DEFAULT_PROCESSED_OUTPUT = PROJECT_ROOT / "data_processed" / "cbonds_bond_cards_processed.xlsx"
DEFAULT_QUOTES_OUTPUT = PROJECT_ROOT / "data_processed" / "cbonds_quotes_needed.xlsx"

MONTHS_RU = {
    "января": 1,
    "февраля": 2,
    "марта": 3,
    "апреля": 4,
    "мая": 5,
    "июня": 6,
    "июля": 7,
    "августа": 8,
    "сентября": 9,
    "октября": 10,
    "ноября": 11,
    "декабря": 12,
}

PLACEHOLDER_VALUES = {"", "-", "—", "–", "нет", "none", "nan"}

ISSUE_INFO_FIELD_MAP = {
    "Статус": "status_raw",
    "Заемщик": "borrower_name",
    "Вид долговых обязательств": "debt_type",
    "Валюта эмиссии": "issue_currency",
    "Номинал / минимальный торговый лот": "nominal_value",
    "Номинал (еврооблигации)": "nominal_eurobond_value",
    "Непогашенный номинал": "outstanding_nominal_value",
    "Анонсированный объём": "announced_volume",
    "Объем эмиссии": "issue_volume",
    "Объем в обращении": "circulation_volume",
    "Объем по непогашенному номиналу": "outstanding_volume",
    "Метод расчета НКД": "accrual_method",
    "Рейтинги эмиссии (АКРА / Эксперт РА)": "issue_rating_text",
    "Листинг": "listing_text",
    "Дата включения в ломбардный список": "lombard_list_date",
    "Номер программы облигаций": "bond_program_number",
    "Гос. регистрационный номер": "reg_number",
    "Дата регистрации": "registration_date",
    "ISIN RegS": "parsed_isin",
    "ISIN 144A": "parsed_isin_144a",
    "CUSIP RegS": "cusip_regs",
    "CUSIP 144A": "cusip_144a",
    "Дата окончания размещения": "placement_end_date",
    "Дата начала начисления купонов": "coupon_start_date",
    "Дата погашения": "maturity_date",
    "Ближайшая оферта (call)": "nearest_call_offer_date",
    "Ближайшая оферта (put)": "nearest_put_offer_date",
}

NUMERIC_FIELDS = {
    "nominal_value",
    "nominal_eurobond_value",
    "outstanding_nominal_value",
    "announced_volume",
    "issue_volume",
    "circulation_volume",
    "outstanding_volume",
}

DATE_FIELDS = {
    "lombard_list_date",
    "registration_date",
    "placement_end_date",
    "coupon_start_date",
    "maturity_date",
    "nearest_call_offer_date",
    "nearest_put_offer_date",
}


def _is_placeholder(value: object) -> bool:
    text = clean_text(value).lower()
    return text in PLACEHOLDER_VALUES


def parse_ru_date(value: object) -> pd.Timestamp | pd.NaT:
    text = clean_text(value)
    if not text or _is_placeholder(text):
        return pd.NaT
    if isinstance(value, pd.Timestamp):
        return value.normalize()

    direct = pd.to_datetime(text, errors="coerce", dayfirst=True)
    if pd.notna(direct):
        return pd.Timestamp(direct).normalize()

    text = re.sub(r"^[А-Яа-яЁё]+,\s*", "", text)
    text = text.replace(" г. в 00:00:00", "")
    text = text.replace(" г.", "")
    text = text.replace(" в 00:00:00", "")
    text = text.replace("\n", " ")
    text = re.sub(r"\s+", " ", text).strip().lower()

    direct = pd.to_datetime(text, errors="coerce", dayfirst=True)
    if pd.notna(direct):
        return pd.Timestamp(direct).normalize()

    match = re.search(r"(\d{1,2})\s+([а-яё]+)\s+(\d{4})", text)
    if not match:
        return pd.NaT
    day = int(match.group(1))
    month = MONTHS_RU.get(match.group(2))
    year = int(match.group(3))
    if not month:
        return pd.NaT
    return pd.Timestamp(year=year, month=month, day=day)


def parse_offer_window_start(value: object) -> pd.Timestamp | pd.NaT:
    text = clean_text(value)
    if not text or _is_placeholder(text):
        return pd.NaT
    match = re.search(r"с\s+(\d{2}\.\d{2}\.\d{4})", text)
    if match:
        return parse_ru_date(match.group(1))
    return parse_ru_date(text)


def parse_offer_window_end(value: object) -> pd.Timestamp | pd.NaT:
    text = clean_text(value)
    if not text or _is_placeholder(text):
        return pd.NaT
    match = re.search(r"по\s+(\d{2}\.\d{2}\.\d{4})", text)
    if match:
        return parse_ru_date(match.group(1))
    return pd.NaT


def classify_status(value: object) -> str:
    text = normalize_text(value)
    if not text:
        return "unknown"
    if "дефолт" in text:
        return "default"
    if "погаш" in text:
        return "redeemed"
    if "обращ" in text:
        return "outstanding"
    if "аннул" in text:
        return "cancelled"
    if "размещ" in text:
        return "placed"
    return "other"


def _core_name(value: object) -> str:
    return strip_legal_form(value)


def borrower_match_score(company_name: object, borrower_name: object) -> float:
    company_core = _core_name(company_name)
    borrower_core = _core_name(borrower_name)
    if not company_core or not borrower_core:
        return 0.0
    if company_core == borrower_core:
        return 1.0
    if borrower_core in company_core or company_core in borrower_core:
        return 0.9
    company_tokens = set(company_core.split())
    borrower_tokens = set(borrower_core.split())
    if not company_tokens or not borrower_tokens:
        return 0.0
    overlap = len(company_tokens & borrower_tokens)
    union = len(company_tokens | borrower_tokens)
    return overlap / union if union else 0.0


def match_bucket(score: float) -> str:
    if score >= 0.8:
        return "high"
    if score >= 0.4:
        return "medium"
    return "low"


def build_sheet_grid(ws, max_row: int = 100, max_col: int = 9) -> list[list[object]]:
    grid: list[list[object]] = []
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True):
        grid.append(list(row))
    return grid


def grid_value(grid: list[list[object]], row: int, column: int) -> object:
    row_idx = row - 1
    col_idx = column - 1
    if row_idx < 0 or row_idx >= len(grid):
        return None
    row_values = grid[row_idx]
    if col_idx < 0 or col_idx >= len(row_values):
        return None
    return row_values[col_idx]


def extract_issue_info(grid: list[list[object]]) -> dict[str, object]:
    raw_pairs: dict[str, object] = {}
    for row in range(6, 20):
        for label_col, value_col in ((1, 2), (4, 5)):
            label = clean_text(grid_value(grid, row, label_col))
            value = grid_value(grid, row, value_col)
            if label:
                raw_pairs[label] = value

    out: dict[str, object] = {}
    for label, column in ISSUE_INFO_FIELD_MAP.items():
        value = raw_pairs.get(label)
        if column in DATE_FIELDS:
            out[column] = parse_ru_date(value)
        elif column == "parsed_isin":
            out[column] = normalize_isin(value)
        elif column in NUMERIC_FIELDS:
            out[column] = to_float(value)
        else:
            out[column] = clean_text(value)
    return out


def find_section_row(grid: list[list[object]], title: str) -> int | None:
    for row in range(1, min(len(grid), 150) + 1):
        if clean_text(grid_value(grid, row, 1)) == title:
            return row
    return None


def parse_coupon_schedule(grid: list[list[object]], sheet_name: str) -> pd.DataFrame:
    start_row = find_section_row(grid, "График выплат")
    if start_row is None:
        return pd.DataFrame()

    rows: list[dict[str, object]] = []
    row = start_row + 2
    while row <= len(grid):
        first_cell = clean_text(grid_value(grid, row, 1))
        second_cell = clean_text(grid_value(grid, row, 2))
        if not first_cell and not second_cell:
            break
        if first_cell == "Условия досрочного выкупа":
            break
        if first_cell in {"№", "Дата"}:
            row += 1
            continue
        rows.append(
            {
                "sheet_name": sheet_name,
                "coupon_number": to_float(first_cell),
                "coupon_end_date": parse_ru_date(grid_value(grid, row, 2)),
                "coupon_payment_date": parse_ru_date(grid_value(grid, row, 3)),
                "coupon_record_date": parse_ru_date(grid_value(grid, row, 4)),
                "coupon_rate_pct": to_float(grid_value(grid, row, 5)),
                "coupon_amount": to_float(grid_value(grid, row, 6)),
                "principal_repayment": to_float(grid_value(grid, row, 7)),
            }
        )
        row += 1
    return pd.DataFrame(rows)


def parse_offer_terms(grid: list[list[object]], sheet_name: str) -> pd.DataFrame:
    start_row = find_section_row(grid, "Условия досрочного выкупа")
    if start_row is None:
        return pd.DataFrame()

    rows: list[dict[str, object]] = []
    row = start_row + 2
    blank_streak = 0
    while row <= len(grid):
        values = [grid_value(grid, row, col) for col in range(1, 10)]
        texts = [clean_text(value) for value in values]
        if not any(texts):
            blank_streak += 1
            if blank_streak >= 2:
                break
            row += 1
            continue
        blank_streak = 0
        if texts[0] in {"Дата", "№"}:
            row += 1
            continue
        rows.append(
            {
                "sheet_name": sheet_name,
                "offer_date": parse_ru_date(values[0]),
                "offer_window_start": parse_offer_window_start(values[1]),
                "offer_window_end": parse_offer_window_end(values[1]),
                "offer_window_raw": texts[1],
                "offer_type": texts[2],
                "option_type": texts[3],
                "benchmark_spread_bp": to_float(values[4]),
                "offer_valid_until": parse_ru_date(values[5]),
                "offer_price_pct": to_float(values[6]),
                "repurchased_volume_mln": to_float(values[7]),
                "offer_notes": texts[8] if len(texts) > 8 else "",
            }
        )
        row += 1
    return pd.DataFrame(rows)


def parse_quote_snapshot(grid: list[list[object]], section_title: str, sheet_name: str) -> pd.DataFrame:
    start_row = find_section_row(grid, section_title)
    if start_row is None:
        return pd.DataFrame()
    row = start_row + 2
    rows: list[dict[str, object]] = []
    while row <= len(grid):
        venue = clean_text(grid_value(grid, row, 1))
        date_text = clean_text(grid_value(grid, row, 2))
        if not venue and not date_text:
            break
        if venue in {"Торговая площадка", "№", "Дата"}:
            row += 1
            continue
        rows.append(
            {
                "sheet_name": sheet_name,
                "section_title": section_title,
                "quote_venue": venue,
                "quote_datetime": parse_ru_date(date_text),
                "quote_bid_pct": to_float(grid_value(grid, row, 3)),
                "quote_ask_pct": to_float(grid_value(grid, row, 4)),
                "quote_yield_bid_pct": to_float(grid_value(grid, row, 5)),
                "quote_yield_ask_pct": to_float(grid_value(grid, row, 6)),
                "quote_indic_price_pct": to_float(grid_value(grid, row, 7)),
                "quote_indic_yield_pct": to_float(grid_value(grid, row, 8)),
            }
        )
        row += 1
    return pd.DataFrame(rows)


def build_issue_record(ws, mapped_row: pd.Series) -> tuple[dict[str, object], pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    grid = build_sheet_grid(ws)
    sheet_name = clean_text(mapped_row.get("sheet_name"))
    issue_name = clean_text(grid_value(grid, 3, 2))
    info = extract_issue_info(grid)
    coupon_schedule = parse_coupon_schedule(grid, sheet_name)
    offer_terms = parse_offer_terms(grid, sheet_name)
    quote_snapshot = pd.concat(
        [
            parse_quote_snapshot(grid, "Cbonds Valuation", sheet_name),
            parse_quote_snapshot(grid, "Биржевые и внебиржевые котировки", sheet_name),
        ],
        ignore_index=True,
        sort=False,
    )

    borrower_score = borrower_match_score(mapped_row.get("company_name"), info.get("borrower_name"))
    parsed_isin = normalize_isin(info.get("parsed_isin")) or normalize_isin(mapped_row.get("isin")) or normalize_isin(mapped_row.get("query_text"))
    issue_key = ""
    if parsed_isin:
        issue_key = f"isin:{parsed_isin}"
    elif clean_text(info.get("reg_number")):
        issue_key = f"reg:{clean_text(info.get('reg_number'))}"
    elif issue_name:
        issue_key = f"name:{normalize_text(issue_name)}"
    else:
        issue_key = f"sheet:{sheet_name}"

    if not coupon_schedule.empty and "coupon_rate_pct" in coupon_schedule.columns:
        coupon_rates = pd.to_numeric(coupon_schedule["coupon_rate_pct"], errors="coerce")
    else:
        coupon_rates = pd.Series(dtype=float)

    if not quote_snapshot.empty and "quote_datetime" in quote_snapshot.columns:
        quote_dates = pd.to_datetime(quote_snapshot["quote_datetime"], errors="coerce")
    else:
        quote_dates = pd.Series(dtype="datetime64[ns]")

    if not offer_terms.empty and "offer_date" in offer_terms.columns:
        offers = pd.to_datetime(offer_terms["offer_date"], errors="coerce")
    else:
        offers = pd.Series(dtype="datetime64[ns]")

    issue_row = {
        "sheet_name": sheet_name,
        "issue_name": issue_name,
        "issue_key": issue_key,
        "query_text": clean_text(mapped_row.get("query_text")),
        "capture_status": clean_text(mapped_row.get("status")),
        "capture_timestamp": clean_text(mapped_row.get("timestamp")),
        "source_row": pd.to_numeric(mapped_row.get("source_row"), errors="coerce"),
        "company_id": clean_text(mapped_row.get("company_id")),
        "sector": clean_text(mapped_row.get("sector")),
        "sample_flag": clean_text(mapped_row.get("sample_flag")),
        "company_name": clean_text(mapped_row.get("company_name")),
        "company_inn": normalize_identifier(mapped_row.get("company_inn"), length=10),
        "company_ticker": clean_text(mapped_row.get("company_ticker")),
        "company_isin": normalize_isin(mapped_row.get("company_isin")),
        "security_ticker": clean_text(mapped_row.get("ticker")),
        "security_isin": normalize_isin(mapped_row.get("isin")),
        "instrument_name": clean_text(mapped_row.get("instrument_name")),
        "best_source": clean_text(mapped_row.get("best_source")),
        "best_source_instrument_id": clean_text(mapped_row.get("best_source_instrument_id")),
        "best_n_trade_dates": pd.to_numeric(mapped_row.get("best_n_trade_dates"), errors="coerce"),
        "best_history_start": pd.to_datetime(mapped_row.get("best_history_start"), errors="coerce"),
        "best_history_end": pd.to_datetime(mapped_row.get("best_history_end"), errors="coerce"),
        "best_match_score": pd.to_numeric(mapped_row.get("best_match_score"), errors="coerce"),
        "best_match_confidence": clean_text(mapped_row.get("best_match_confidence")),
        "manual_review_needed_flag": bool(mapped_row.get("manual_review_needed_flag", False)),
        "history_available_flag": bool(mapped_row.get("history_available_flag", False)),
        "usable_history_flag": bool(mapped_row.get("usable_history_flag", False)),
        "reliable_mapping_flag": bool(mapped_row.get("reliable_mapping_flag", False)),
        "market_relevance_bucket": clean_text(mapped_row.get("market_relevance_bucket")),
        **info,
        "status_category": classify_status(info.get("status_raw")),
        "borrower_match_score": borrower_score,
        "borrower_match_bucket": match_bucket(borrower_score),
        "borrower_match_flag": borrower_score >= 0.4,
        "issue_rating_present_flag": not _is_placeholder(info.get("issue_rating_text")),
        "listing_present_flag": not _is_placeholder(info.get("listing_text")),
        "listing_delisted_flag": "делист" in normalize_text(info.get("listing_text")),
        "has_offer_flag": not offer_terms.empty or pd.notna(info.get("nearest_call_offer_date")) or pd.notna(info.get("nearest_put_offer_date")),
        "coupon_rows": int(len(coupon_schedule)),
        "offer_rows": int(len(offer_terms)),
        "quote_snapshot_rows": int(len(quote_snapshot)),
        "coupon_rate_first_pct": coupon_rates.iloc[0] if len(coupon_rates.dropna()) else np.nan,
        "coupon_rate_last_pct": coupon_rates.iloc[-1] if len(coupon_rates.dropna()) else np.nan,
        "coupon_rate_avg_pct": coupon_rates.mean() if len(coupon_rates.dropna()) else np.nan,
        "coupon_rate_max_pct": coupon_rates.max() if len(coupon_rates.dropna()) else np.nan,
        "coupon_first_end_date": coupon_schedule["coupon_end_date"].min() if "coupon_end_date" in coupon_schedule.columns and not coupon_schedule.empty else pd.NaT,
        "coupon_last_end_date": coupon_schedule["coupon_end_date"].max() if "coupon_end_date" in coupon_schedule.columns and not coupon_schedule.empty else pd.NaT,
        "offer_first_date": offers.min() if len(offers.dropna()) else pd.NaT,
        "offer_last_date": offers.max() if len(offers.dropna()) else pd.NaT,
        "quote_snapshot_latest_date": quote_dates.max() if len(quote_dates.dropna()) else pd.NaT,
        "quote_snapshot_venue_count": int(quote_snapshot["quote_venue"].replace("", np.nan).dropna().nunique()) if not quote_snapshot.empty else 0,
    }
    return issue_row, coupon_schedule, offer_terms, quote_snapshot


def dedupe_issue_master(issue_master_raw: pd.DataFrame) -> pd.DataFrame:
    if issue_master_raw.empty:
        return issue_master_raw
    out = issue_master_raw.copy()
    out["dedupe_score"] = (
        out["reliable_mapping_flag"].fillna(False).astype(int) * 100
        + out["borrower_match_flag"].fillna(False).astype(int) * 80
        + out["usable_history_flag"].fillna(False).astype(int) * 20
        + out["history_available_flag"].fillna(False).astype(int) * 10
        + out["sample_flag"].eq("main").astype(int) * 10
        - out["manual_review_needed_flag"].fillna(False).astype(int) * 40
    )
    out = out.sort_values(
        [
            "issue_key",
            "dedupe_score",
            "best_n_trade_dates",
            "borrower_match_score",
        ],
        ascending=[True, False, False, False],
        kind="stable",
    )
    deduped = out.drop_duplicates(subset=["issue_key"], keep="first").copy()
    return deduped.drop(columns=["dedupe_score"], errors="ignore").reset_index(drop=True)


def add_quotes_priority(issue_master: pd.DataFrame) -> pd.DataFrame:
    if issue_master.empty:
        return issue_master
    out = issue_master.copy()
    high_conf = (
        out["reliable_mapping_flag"].fillna(False)
        | out["borrower_match_flag"].fillna(False)
    ) & (~out["manual_review_needed_flag"].fillna(False))

    lower_need = high_conf & out["usable_history_flag"].fillna(False) & out["best_n_trade_dates"].fillna(0).ge(750)

    out["quotes_priority_bucket"] = np.select(
        [
            ~high_conf,
            lower_need,
            high_conf,
        ],
        [
            "manual_review_first",
            "needed_lower_priority",
            "definitely_needed",
        ],
        default="needed_lower_priority",
    )
    out["quotes_priority_score"] = (
        out["sample_flag"].eq("main").astype(int) * 20
        + out["reliable_mapping_flag"].fillna(False).astype(int) * 40
        + out["borrower_match_flag"].fillna(False).astype(int) * 30
        + (~out["usable_history_flag"].fillna(False)).astype(int) * 25
        + (~out["history_available_flag"].fillna(False)).astype(int) * 10
        + out["status_category"].isin(["outstanding", "redeemed", "default"]).astype(int) * 5
        - out["manual_review_needed_flag"].fillna(False).astype(int) * 30
    )
    out["quotes_needed_flag"] = out["quotes_priority_bucket"].isin(["definitely_needed", "needed_lower_priority"])
    return out


def add_model_use_flag(issue_master: pd.DataFrame) -> pd.DataFrame:
    if issue_master.empty:
        return issue_master
    out = issue_master.copy()
    out["model_use_flag"] = (
        (
            out["reliable_mapping_flag"].fillna(False)
            | pd.to_numeric(out["borrower_match_score"], errors="coerce").fillna(0).ge(0.8)
        )
        & (~out["manual_review_needed_flag"].fillna(False))
    )
    return out


def build_company_bond_summary(issue_master: pd.DataFrame) -> pd.DataFrame:
    if issue_master.empty:
        return pd.DataFrame()

    def _sum_if(frame: pd.DataFrame, mask: pd.Series, column: str) -> float:
        series = pd.to_numeric(frame.loc[mask, column], errors="coerce")
        return float(series.sum(min_count=1)) if not series.dropna().empty else np.nan

    rows: list[dict[str, object]] = []
    for company_id, frame in issue_master.groupby("company_id", dropna=False):
        issue_count = len(frame)
        reliable = frame["reliable_mapping_flag"].fillna(False)
        borrower_ok = frame["borrower_match_flag"].fillna(False)
        high_conf = (reliable | borrower_ok) & (~frame["manual_review_needed_flag"].fillna(False))
        outst = frame["status_category"].eq("outstanding")
        rows.append(
            {
                "company_id": company_id,
                "cbonds_issue_card_found_flag": True,
                "cbonds_issue_count": issue_count,
                "cbonds_issue_count_high_conf": int(high_conf.sum()),
                "cbonds_issue_count_outstanding": int(outst.sum()),
                "cbonds_issue_count_redeemed": int(frame["status_category"].eq("redeemed").sum()),
                "cbonds_issue_count_default": int(frame["status_category"].eq("default").sum()),
                "cbonds_issue_count_need_quotes_high": int(frame["quotes_priority_bucket"].eq("definitely_needed").sum()),
                "cbonds_issue_count_need_quotes_lower": int(frame["quotes_priority_bucket"].eq("needed_lower_priority").sum()),
                "cbonds_issue_count_manual_review": int(frame["quotes_priority_bucket"].eq("manual_review_first").sum()),
                "cbonds_issue_currency_count": int(frame["issue_currency"].replace("", np.nan).dropna().nunique()),
                "cbonds_issue_volume_sum": _sum_if(frame, pd.Series(True, index=frame.index), "issue_volume"),
                "cbonds_issue_volume_rub_sum": _sum_if(frame, frame["issue_currency"].eq("RUB"), "issue_volume"),
                "cbonds_circulation_volume_sum": _sum_if(frame, pd.Series(True, index=frame.index), "circulation_volume"),
                "cbonds_circulation_volume_rub_sum": _sum_if(frame, frame["issue_currency"].eq("RUB"), "circulation_volume"),
                "cbonds_outstanding_issue_volume_rub_sum": _sum_if(frame, outst & frame["issue_currency"].eq("RUB"), "circulation_volume"),
                "cbonds_has_rating_issue_flag": bool(frame["issue_rating_present_flag"].fillna(False).any()),
                "cbonds_has_listing_issue_flag": bool(frame["listing_present_flag"].fillna(False).any()),
                "cbonds_has_offer_issue_flag": bool(frame["has_offer_flag"].fillna(False).any()),
                "cbonds_first_registration_date": pd.to_datetime(frame["registration_date"], errors="coerce").min(),
                "cbonds_first_placement_end_date": pd.to_datetime(frame["placement_end_date"], errors="coerce").min(),
                "cbonds_last_maturity_date": pd.to_datetime(frame["maturity_date"], errors="coerce").max(),
                "cbonds_coupon_rate_avg_pct": pd.to_numeric(frame["coupon_rate_avg_pct"], errors="coerce").mean(),
                "cbonds_coupon_rate_max_pct": pd.to_numeric(frame["coupon_rate_max_pct"], errors="coerce").max(),
                "cbonds_borrower_match_high_conf_flag": bool(frame["borrower_match_bucket"].eq("high").any()),
            }
        )
    return pd.DataFrame(rows).sort_values("company_id", kind="stable").reset_index(drop=True)


def _quarter_label(series: pd.Series) -> pd.Series:
    dt = pd.to_datetime(series, errors="coerce")
    q = ((dt.dt.month - 1) // 3 + 1).astype("Int64")
    return dt.dt.year.astype("Int64").astype(str) + "Q" + q.astype(str)


def build_company_bond_quarterly(issue_master: pd.DataFrame, offer_terms: pd.DataFrame) -> pd.DataFrame:
    if issue_master.empty:
        return pd.DataFrame()

    events: list[pd.DataFrame] = []

    for event_name, date_col, volume_col in [
        ("cbonds_issue_registration_q", "registration_date", "issue_volume"),
        ("cbonds_issue_placement_q", "placement_end_date", "issue_volume"),
        ("cbonds_issue_maturity_q", "maturity_date", "circulation_volume"),
    ]:
        subset = issue_master[["company_id", date_col, volume_col]].copy()
        subset["quarter_label"] = _quarter_label(subset[date_col])
        subset = subset.loc[subset["quarter_label"].str.contains("Q", na=False)].copy()
        if subset.empty:
            continue
        agg = (
            subset.groupby(["company_id", "quarter_label"], dropna=False)
            .agg(
                **{
                    f"{event_name}_count": (date_col, "size"),
                    f"{event_name}_volume_rub": (
                        volume_col,
                        lambda values: pd.to_numeric(values, errors="coerce").sum(min_count=1),
                    ),
                }
            )
            .reset_index()
        )
        events.append(agg)

    if not offer_terms.empty:
        offers = offer_terms.copy()
        offers["quarter_label"] = _quarter_label(offers["offer_date"])
        offers = offers.loc[offers["quarter_label"].str.contains("Q", na=False)].copy()
        if not offers.empty:
            events.append(
                offers.groupby(["company_id", "quarter_label"], dropna=False)
                .agg(
                    cbonds_issue_offer_q_count=("offer_date", "size"),
                    cbonds_issue_offer_q_price_avg=("offer_price_pct", "mean"),
                )
                .reset_index()
            )

    outstanding_rows: list[dict[str, object]] = []
    quarter_grid = pd.period_range("2014Q3", "2025Q4", freq="Q")
    for _, row in issue_master.iterrows():
        start = pd.to_datetime(row.get("placement_end_date"), errors="coerce")
        if pd.isna(start):
            start = pd.to_datetime(row.get("registration_date"), errors="coerce")
        end = pd.to_datetime(row.get("maturity_date"), errors="coerce")
        if pd.isna(start) or pd.isna(end):
            continue
        start_period = pd.Period(start, freq="Q")
        end_period = pd.Period(end, freq="Q")
        active_periods = quarter_grid[(quarter_grid >= start_period) & (quarter_grid <= end_period)]
        volume_rub = row.get("circulation_volume") if clean_text(row.get("issue_currency")) == "RUB" else np.nan
        for period in active_periods:
            outstanding_rows.append(
                {
                    "company_id": row["company_id"],
                    "quarter_label": f"{period.year}Q{period.quarter}",
                    "cbonds_issue_outstanding_q_count": 1,
                    "cbonds_issue_outstanding_q_volume_rub": volume_rub,
                }
            )
    if outstanding_rows:
        outstanding = (
            pd.DataFrame(outstanding_rows)
            .groupby(["company_id", "quarter_label"], dropna=False)
            .agg(
                cbonds_issue_outstanding_q_count=("cbonds_issue_outstanding_q_count", "sum"),
                cbonds_issue_outstanding_q_volume_rub=("cbonds_issue_outstanding_q_volume_rub", lambda values: pd.to_numeric(values, errors="coerce").sum(min_count=1)),
            )
            .reset_index()
        )
        events.append(outstanding)

    if not events:
        return pd.DataFrame()

    out = events[0]
    for frame in events[1:]:
        out = out.merge(frame, on=["company_id", "quarter_label"], how="outer")
    return out.sort_values(["company_id", "quarter_label"], kind="stable").reset_index(drop=True)


def parse_capture_buckets(capture_paths: list[Path], bond_queries: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    bond_queries = bond_queries.copy().reset_index().rename(columns={"index": "idx0"})

    issue_rows: list[dict[str, object]] = []
    coupon_frames: list[pd.DataFrame] = []
    offer_frames: list[pd.DataFrame] = []
    quote_frames: list[pd.DataFrame] = []
    logs: list[dict[str, object]] = []

    for workbook_path in capture_paths:
        capture_log = pd.read_excel(workbook_path, sheet_name="capture_log")
        capture_log["idx0"] = pd.to_numeric(capture_log["source_row"], errors="coerce") - 2
        mapped = capture_log.merge(bond_queries, on="idx0", how="left")
        wb = load_workbook(workbook_path, read_only=True, data_only=True)

        for _, row in mapped.iterrows():
            sheet_name = clean_text(row.get("sheet_name"))
            if not sheet_name or sheet_name not in wb.sheetnames:
                logs.append(
                    {
                        "workbook": workbook_path.name,
                        "sheet_name": sheet_name,
                        "query_text": clean_text(row.get("query_text")),
                        "company_id": clean_text(row.get("company_id")),
                        "parse_status": "sheet_missing",
                    }
                )
                continue
            ws = wb[sheet_name]
            try:
                issue_row, coupon_schedule, offer_terms, quote_snapshot = build_issue_record(ws, row)
                issue_row["source_workbook"] = workbook_path.name
                issue_rows.append(issue_row)

                if not coupon_schedule.empty:
                    coupon_schedule["source_workbook"] = workbook_path.name
                    coupon_schedule["issue_key"] = issue_row["issue_key"]
                    coupon_schedule["company_id"] = issue_row["company_id"]
                    coupon_schedule["issue_name"] = issue_row["issue_name"]
                    coupon_schedule["parsed_isin"] = issue_row["parsed_isin"]
                    coupon_frames.append(coupon_schedule)

                if not offer_terms.empty:
                    offer_terms["source_workbook"] = workbook_path.name
                    offer_terms["issue_key"] = issue_row["issue_key"]
                    offer_terms["company_id"] = issue_row["company_id"]
                    offer_terms["issue_name"] = issue_row["issue_name"]
                    offer_terms["parsed_isin"] = issue_row["parsed_isin"]
                    offer_frames.append(offer_terms)

                if not quote_snapshot.empty:
                    quote_snapshot["source_workbook"] = workbook_path.name
                    quote_snapshot["issue_key"] = issue_row["issue_key"]
                    quote_snapshot["company_id"] = issue_row["company_id"]
                    quote_snapshot["issue_name"] = issue_row["issue_name"]
                    quote_snapshot["parsed_isin"] = issue_row["parsed_isin"]
                    quote_frames.append(quote_snapshot)

                logs.append(
                    {
                        "workbook": workbook_path.name,
                        "sheet_name": sheet_name,
                        "query_text": clean_text(row.get("query_text")),
                        "company_id": clean_text(row.get("company_id")),
                        "parse_status": "ok",
                    }
                )
            except Exception as exc:  # pragma: no cover - defensive logging
                logs.append(
                    {
                        "workbook": workbook_path.name,
                        "sheet_name": sheet_name,
                        "query_text": clean_text(row.get("query_text")),
                        "company_id": clean_text(row.get("company_id")),
                        "parse_status": "error",
                        "error": str(exc),
                    }
                )

    issue_master_raw = pd.DataFrame(issue_rows)
    coupon_schedule = pd.concat(coupon_frames, ignore_index=True, sort=False) if coupon_frames else pd.DataFrame()
    offer_terms = pd.concat(offer_frames, ignore_index=True, sort=False) if offer_frames else pd.DataFrame()
    quote_snapshot = pd.concat(quote_frames, ignore_index=True, sort=False) if quote_frames else pd.DataFrame()
    parse_log = pd.DataFrame(logs)
    return issue_master_raw, coupon_schedule, offer_terms, quote_snapshot, parse_log


def build_quotes_needed(issue_master: pd.DataFrame) -> pd.DataFrame:
    if issue_master.empty:
        return issue_master
    columns = [
        "company_id",
        "sector",
        "sample_flag",
        "company_name",
        "company_inn",
        "issue_name",
        "parsed_isin",
        "reg_number",
        "issue_currency",
        "status_raw",
        "status_category",
        "issue_volume",
        "circulation_volume",
        "registration_date",
        "placement_end_date",
        "maturity_date",
        "issue_rating_text",
        "listing_text",
        "coupon_rate_last_pct",
        "coupon_rate_max_pct",
        "offer_rows",
        "has_offer_flag",
        "best_source",
        "best_n_trade_dates",
        "history_available_flag",
        "usable_history_flag",
        "reliable_mapping_flag",
        "borrower_name",
        "borrower_match_score",
        "borrower_match_bucket",
        "manual_review_needed_flag",
        "quotes_priority_bucket",
        "quotes_priority_score",
        "quotes_needed_flag",
    ]
    out = issue_master[columns].copy()
    out["cbonds_quote_query_primary"] = out["parsed_isin"].where(out["parsed_isin"].astype(str).str.len().gt(0), out["issue_name"])
    out["cbonds_quote_query_secondary"] = out["issue_name"]
    out = out.sort_values(
        ["quotes_needed_flag", "quotes_priority_score", "sample_flag", "company_name", "issue_name"],
        ascending=[False, False, True, True, True],
        kind="stable",
    ).reset_index(drop=True)
    return out


def build_summary(issue_master: pd.DataFrame, quotes_needed: pd.DataFrame, parse_log: pd.DataFrame) -> pd.DataFrame:
    if issue_master.empty:
        return pd.DataFrame([{"metric": "issue_cards_parsed", "value": 0}])
    rows = [
        {"metric": "issue_cards_parsed", "value": len(issue_master)},
        {"metric": "unique_issues_clean", "value": int(issue_master["issue_key"].nunique())},
        {"metric": "companies_with_bonds_found", "value": int(issue_master["company_id"].nunique())},
        {"metric": "companies_with_bonds_model_use", "value": int(issue_master.loc[issue_master["model_use_flag"].fillna(False), "company_id"].nunique())},
        {"metric": "companies_oil_gas_with_bonds_found", "value": int(issue_master.loc[issue_master["sector"].eq("oil_gas"), "company_id"].nunique())},
        {"metric": "companies_metallurgy_with_bonds_found", "value": int(issue_master.loc[issue_master["sector"].eq("metallurgy"), "company_id"].nunique())},
        {"metric": "issues_model_use", "value": int(issue_master["model_use_flag"].fillna(False).sum())},
        {"metric": "issues_definitely_need_quotes", "value": int(quotes_needed["quotes_priority_bucket"].eq("definitely_needed").sum())},
        {"metric": "issues_need_quotes_lower_priority", "value": int(quotes_needed["quotes_priority_bucket"].eq("needed_lower_priority").sum())},
        {"metric": "issues_manual_review_first", "value": int(quotes_needed["quotes_priority_bucket"].eq("manual_review_first").sum())},
        {"metric": "issues_with_rating", "value": int(issue_master["issue_rating_present_flag"].fillna(False).sum())},
        {"metric": "issues_with_offer", "value": int(issue_master["has_offer_flag"].fillna(False).sum())},
        {"metric": "issues_with_listing", "value": int(issue_master["listing_present_flag"].fillna(False).sum())},
        {"metric": "parse_errors", "value": int(parse_log["parse_status"].eq("error").sum()) if not parse_log.empty else 0},
    ]
    return pd.DataFrame(rows)


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    ensure_directory(path.parent)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)


def build_cbonds_bond_cards_outputs(
    helper_workbook: Path = DEFAULT_HELPER_WORKBOOK,
    capture_dir: Path = PROJECT_ROOT / "data_processed",
    capture_glob: str = DEFAULT_CAPTURE_GLOB,
    processed_output: Path = DEFAULT_PROCESSED_OUTPUT,
    quotes_output: Path = DEFAULT_QUOTES_OUTPUT,
) -> dict[str, pd.DataFrame]:
    bond_queries = pd.read_excel(helper_workbook, sheet_name="bond_queries")
    capture_paths = sorted(capture_dir.glob(capture_glob))
    if not capture_paths:
        raise FileNotFoundError(f"No Cbonds capture bucket files found by pattern {capture_glob!r} in {capture_dir}")

    issue_master_raw, coupon_schedule, offer_terms, quote_snapshot, parse_log = parse_capture_buckets(capture_paths, bond_queries)
    issue_master_clean = add_model_use_flag(add_quotes_priority(dedupe_issue_master(issue_master_raw)))
    issue_master_model = issue_master_clean.loc[issue_master_clean["model_use_flag"].fillna(False)].copy()

    if not offer_terms.empty:
        offer_terms = offer_terms.merge(
            issue_master_clean[["sheet_name", "company_id", "issue_key", "parsed_isin", "issue_name"]],
            on=["sheet_name", "company_id", "issue_key", "parsed_isin", "issue_name"],
            how="left",
        )
    offer_terms_model = offer_terms.loc[offer_terms["issue_key"].isin(issue_master_model["issue_key"])].copy() if not offer_terms.empty else pd.DataFrame()

    company_bond_summary = build_company_bond_summary(issue_master_clean)
    company_bond_quarterly = build_company_bond_quarterly(issue_master_clean, offer_terms)
    company_bond_summary_model = build_company_bond_summary(issue_master_model)
    company_bond_quarterly_model = build_company_bond_quarterly(issue_master_model, offer_terms_model)
    quotes_needed = build_quotes_needed(issue_master_clean)
    summary = build_summary(issue_master_clean, quotes_needed, parse_log)

    processed_sheets = {
        "summary": summary,
        "issue_master_raw": issue_master_raw,
        "issue_master_clean": issue_master_clean,
        "issue_master_model": issue_master_model,
        "company_bond_summary": company_bond_summary,
        "company_bond_summary_model": company_bond_summary_model,
        "company_bond_quarterly": company_bond_quarterly,
        "company_bond_quarterly_model": company_bond_quarterly_model,
        "coupon_schedule": coupon_schedule,
        "offer_terms": offer_terms,
        "offer_terms_model": offer_terms_model,
        "quote_snapshot": quote_snapshot,
        "parse_log": parse_log,
    }
    write_workbook(processed_output, processed_sheets)

    quotes_sheets = {
        "summary": summary,
        "quotes_needed": quotes_needed,
        "issue_master_clean": issue_master_clean,
        "issue_master_model": issue_master_model,
        "company_bond_summary": company_bond_summary,
        "company_bond_summary_model": company_bond_summary_model,
        "manual_review_first": quotes_needed.loc[quotes_needed["quotes_priority_bucket"].eq("manual_review_first")].copy(),
    }
    write_workbook(quotes_output, quotes_sheets)

    return {
        "summary": summary,
        "issue_master_raw": issue_master_raw,
        "issue_master_clean": issue_master_clean,
        "issue_master_model": issue_master_model,
        "company_bond_summary": company_bond_summary,
        "company_bond_summary_model": company_bond_summary_model,
        "company_bond_quarterly": company_bond_quarterly,
        "company_bond_quarterly_model": company_bond_quarterly_model,
        "coupon_schedule": coupon_schedule,
        "offer_terms": offer_terms,
        "offer_terms_model": offer_terms_model,
        "quote_snapshot": quote_snapshot,
        "parse_log": parse_log,
        "quotes_needed": quotes_needed,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Parse Cbonds issue card capture buckets into normalized bond-card datasets.")
    parser.add_argument("--helper", type=Path, default=DEFAULT_HELPER_WORKBOOK)
    parser.add_argument("--capture-dir", type=Path, default=PROJECT_ROOT / "data_processed")
    parser.add_argument("--capture-glob", default=DEFAULT_CAPTURE_GLOB)
    parser.add_argument("--processed-output", type=Path, default=DEFAULT_PROCESSED_OUTPUT)
    parser.add_argument("--quotes-output", type=Path, default=DEFAULT_QUOTES_OUTPUT)
    args = parser.parse_args()

    outputs = build_cbonds_bond_cards_outputs(
        helper_workbook=args.helper,
        capture_dir=args.capture_dir,
        capture_glob=args.capture_glob,
        processed_output=args.processed_output,
        quotes_output=args.quotes_output,
    )

    print(f"Processed workbook: {args.processed_output}")
    print(f"Quotes-needed workbook: {args.quotes_output}")
    print(f"Unique issues: {len(outputs['issue_master_clean'])}")
    print(f"Companies with bonds found: {outputs['issue_master_clean']['company_id'].nunique() if not outputs['issue_master_clean'].empty else 0}")


if __name__ == "__main__":
    main()
