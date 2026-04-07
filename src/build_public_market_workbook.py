from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from .utils import clean_text, save_dataframe_csv


PUBLIC_WORKBOOK_DEFAULT = Path("data_processed/additional_public_data.xlsx")
VALIDATED_WORKBOOK_DEFAULT = Path("data_processed/additional_public_data_validated.xlsx")
OUTPUT_WORKBOOK_DEFAULT = Path("data_processed/public_market_data_full.xlsx")


def bool_value(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return clean_text(value).lower() in {"true", "1", "yes", "да"}


def read_workbook(path: Path | None) -> dict[str, pd.DataFrame]:
    if path is None or not path.exists():
        return {}
    return pd.read_excel(path, sheet_name=None)


def read_csv_sidecar(processed_dir: Path | None, sheet_name: str) -> pd.DataFrame | None:
    if processed_dir is None:
        return None
    path = processed_dir / f"{sheet_name}.csv"
    if not path.exists():
        return None
    try:
        return pd.read_csv(path)
    except pd.errors.EmptyDataError:
        return pd.DataFrame()


def join_unique(values: pd.Series) -> str:
    seen: set[str] = set()
    result: list[str] = []
    for value in values.dropna().astype(str):
        text = clean_text(value)
        if not text or text in seen:
            continue
        seen.add(text)
        result.append(text)
    return ", ".join(result)


def nonempty_sheet(frame: pd.DataFrame | None) -> bool:
    return frame is not None and not frame.empty


def text_series(frame: pd.DataFrame, column: str) -> pd.Series:
    if column in frame.columns:
        return frame[column].fillna("").astype(str)
    return pd.Series([""] * len(frame), index=frame.index, dtype="object")


def add_count_features(
    base: pd.DataFrame,
    source_frame: pd.DataFrame,
    group_column: str,
    prefix: str,
) -> pd.DataFrame:
    if source_frame.empty or group_column not in source_frame.columns:
        base[f"n_{prefix}_rows"] = 0
        base[f"has_{prefix}"] = False
        return base

    counts = (
        source_frame.groupby(group_column, dropna=False)
        .size()
        .reset_index(name=f"n_{prefix}_rows")
        .rename(columns={group_column: "company_name"})
    )
    out = base.merge(counts, on="company_name", how="left")
    out[f"n_{prefix}_rows"] = out[f"n_{prefix}_rows"].fillna(0).astype(int)
    out[f"has_{prefix}"] = out[f"n_{prefix}_rows"].gt(0)
    return out


def selected_subset(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty or "selected_flag" not in frame.columns:
        return pd.DataFrame()
    return frame.loc[frame["selected_flag"].map(bool_value)].copy()


def build_firm_market_flags(
    companies_master: pd.DataFrame,
    public_sheets: dict[str, pd.DataFrame],
    validated_sheets: dict[str, pd.DataFrame],
) -> pd.DataFrame:
    companies = companies_master.copy()
    if companies.empty:
        return companies

    flags = companies[
        [
            column
            for column in [
                "company_name",
                "inn",
                "ogrn",
                "ticker",
                "isin",
                "industry",
                "sample_flag",
                "sample_membership",
            ]
            if column in companies.columns
        ]
    ].copy()

    validated_mapping = validated_sheets.get("issuer_mapping_validated", pd.DataFrame())
    if nonempty_sheet(validated_mapping) and "company_name" in validated_mapping.columns:
        mapping = validated_mapping.copy()
        mapping["instrument_kind_norm"] = text_series(mapping, "instrument_kind").str.lower()
        mapping["validated_bond_flag"] = (
            mapping["instrument_kind_norm"].str.contains("bond", na=False)
            | text_series(mapping, "instrument_id").str.upper().str.startswith("RU")
            | text_series(mapping, "isin").str.upper().str.startswith("RU000A")
        )
        mapping["validated_share_flag"] = mapping["instrument_kind_norm"].str.contains("share", na=False)

        summary = mapping.groupby("company_name", dropna=False).agg(
            n_validated_instruments=("instrument_id", "size"),
            n_validated_bonds=("validated_bond_flag", "sum"),
            n_validated_shares=("validated_share_flag", "sum"),
            validated_sources=("instrument_source", join_unique),
            validated_instrument_ids=("instrument_id", join_unique),
            validated_isins=("isin", join_unique),
        ).reset_index()
        flags = flags.merge(summary, on="company_name", how="left")
    else:
        for column in [
            "n_validated_instruments",
            "n_validated_bonds",
            "n_validated_shares",
            "validated_sources",
            "validated_instrument_ids",
            "validated_isins",
        ]:
            flags[column] = 0 if column.startswith("n_") else ""

    for column in ["n_validated_instruments", "n_validated_bonds", "n_validated_shares"]:
        flags[column] = flags[column].fillna(0).astype(int)
    for column in ["validated_sources", "validated_instrument_ids", "validated_isins"]:
        flags[column] = flags[column].fillna("")
    flags["has_validated_public_instrument"] = flags["n_validated_instruments"].gt(0)
    flags["has_validated_public_bond"] = flags["n_validated_bonds"].gt(0)
    flags["has_validated_public_share"] = flags["n_validated_shares"].gt(0)

    moex_selected = selected_subset(public_sheets.get("moex_instruments", pd.DataFrame()))
    tinvest_selected = selected_subset(public_sheets.get("tinvest_instruments", pd.DataFrame()))
    flags = add_count_features(flags, moex_selected, "company_name", "moex_selected")
    flags = add_count_features(flags, tinvest_selected, "company_name", "tinvest_selected")
    flags = add_count_features(flags, public_sheets.get("moex_bonds", pd.DataFrame()), "company_name", "moex_bonds")
    flags = add_count_features(flags, public_sheets.get("tinvest_bonds", pd.DataFrame()), "company_name", "tinvest_bonds")

    return flags.sort_values(["sample_flag", "company_name"], kind="stable").reset_index(drop=True)


def build_public_market_summary(outputs: dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows = []
    for sheet_name, frame in outputs.items():
        rows.append(
            {
                "sheet_name": sheet_name,
                "n_rows": len(frame),
                "n_columns": len(frame.columns),
                "is_empty": frame.empty,
            }
        )
    return pd.DataFrame(rows)


def reorder_outputs(outputs: dict[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    preferred_order = [
        "public_market_summary",
        "companies_master",
        "firm_market_flags",
        "issuer_mapping_validated",
        "validated_company_actions",
        "moex_instruments_validated",
        "moex_bonds_validated",
        "moex_bond_history",
        "moex_share_history",
        "tinvest_instruments_validated",
        "tinvest_bonds_validated",
        "tinvest_bond_coupons",
        "macro_cbr",
        "commodity_prices_monthly",
        "commodity_prices_annual",
        "mapping_log",
        "download_log",
        "download_log_moex_share_history",
        "manual_validation_checks",
        "cbonds_raw",
    ]
    ordered = {sheet: outputs[sheet] for sheet in preferred_order if sheet in outputs}
    for sheet_name, frame in outputs.items():
        if sheet_name not in ordered:
            ordered[sheet_name] = frame
    return ordered


def format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        ws.freeze_panes = "A2"
        for header_cell in ws[1]:
            header_cell.font = Font(bold=True)
            header_cell.alignment = Alignment(vertical="top", wrap_text=True)
        if ws.max_row <= 5_000 and ws.max_column <= 60:
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_letter in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"]:
            ws.column_dimensions[column_letter].width = 24
    workbook.save(path)


def build_public_market_workbook(
    public_workbook: Path = PUBLIC_WORKBOOK_DEFAULT,
    validated_workbook: Path | None = VALIDATED_WORKBOOK_DEFAULT,
    output_workbook: Path = OUTPUT_WORKBOOK_DEFAULT,
    processed_dir: Path | None = None,
) -> dict[str, pd.DataFrame]:
    public_sheets = read_workbook(public_workbook)
    if not public_sheets:
        raise FileNotFoundError(f"Public workbook not found or empty: {public_workbook}")

    validated_sheets = read_workbook(validated_workbook)
    companies_master = public_sheets.get("companies_master", pd.DataFrame())
    firm_market_flags = build_firm_market_flags(companies_master, public_sheets, validated_sheets)

    outputs: dict[str, pd.DataFrame] = {
        "companies_master": companies_master,
        "firm_market_flags": firm_market_flags,
    }

    if validated_sheets:
        for sheet_name in [
            "issuer_mapping_validated",
            "validated_company_actions",
            "moex_instruments_validated",
            "moex_bonds_validated",
            "tinvest_instruments_validated",
            "tinvest_bonds_validated",
            "manual_validation_checks",
        ]:
            if sheet_name in validated_sheets:
                outputs[sheet_name] = validated_sheets[sheet_name]
    else:
        for sheet_name in ["moex_instruments", "moex_bonds", "tinvest_instruments", "tinvest_bonds"]:
            if sheet_name in public_sheets:
                outputs[sheet_name] = public_sheets[sheet_name]

    for sheet_name in [
        "moex_bond_history",
        "moex_share_history",
        "tinvest_bond_coupons",
        "macro_cbr",
        "commodity_prices_monthly",
        "commodity_prices_annual",
        "mapping_log",
        "download_log",
        "cbonds_raw",
    ]:
        if sheet_name in public_sheets:
            outputs[sheet_name] = public_sheets[sheet_name]
        else:
            sidecar = read_csv_sidecar(processed_dir, sheet_name)
            if sidecar is not None:
                outputs[sheet_name] = sidecar

    share_history_log = read_csv_sidecar(processed_dir, "download_log_moex_share_history")
    if share_history_log is not None and not share_history_log.empty:
        outputs["download_log"] = pd.concat(
            [outputs.get("download_log", pd.DataFrame()), share_history_log],
            ignore_index=True,
            sort=False,
        )
        outputs["download_log_moex_share_history"] = share_history_log

    outputs = reorder_outputs(outputs)
    outputs["public_market_summary"] = build_public_market_summary(outputs)
    outputs = reorder_outputs(outputs)

    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_workbook, engine="openpyxl") as writer:
        for sheet_name, frame in outputs.items():
            frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)

    format_workbook(output_workbook)

    if processed_dir is not None:
        processed_dir.mkdir(parents=True, exist_ok=True)
        save_dataframe_csv(firm_market_flags, processed_dir / "firm_market_flags.csv")

    return outputs
