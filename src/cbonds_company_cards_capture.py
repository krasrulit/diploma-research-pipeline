from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE_CSV = PROJECT_ROOT / "data_processed" / "ownership_state_table.csv"
DEFAULT_OUTPUT_WORKBOOK = PROJECT_ROOT / "data_processed" / "cbonds_company_cards_test_nornickel.xlsx"
DEFAULT_UI_CONFIG = PROJECT_ROOT / "data_processed" / "cbonds_company_cards_ui.json"

ROW_SEPARATOR = chr(30)
CELL_SEPARATOR = chr(31)
EMPTY_MARKER = chr(29)


ACCESSIBILITY_CHECK_SCRIPT = """
tell application "System Events"
\treturn (UI elements enabled as string)
end tell
"""


WORKBOOK_SHEET_CHECK_SCRIPT = """
on run argv
\tset workbookName to item 1 of argv
\tset worksheetName to item 2 of argv
\ttell application "Microsoft Excel"
\t\tget name of worksheet (worksheetName as string) of workbook (workbookName as string)
\tend tell
\treturn "ok"
end run
"""


LIST_WORKSHEETS_SCRIPT = """
on run argv
\tset workbookName to item 1 of argv
\ttell application "Microsoft Excel"
\t\tset worksheetNames to get name of every worksheet of workbook (workbookName as string)
\tend tell
\tset oldTID to AppleScript's text item delimiters
\tset AppleScript's text item delimiters to linefeed
\tset outputText to worksheetNames as text
\tset AppleScript's text item delimiters to oldTID
\treturn outputText
end run
"""


WINDOW_GEOMETRY_SCRIPT = """
tell application "Microsoft Excel" to activate
delay 0.3
tell application "System Events"
\ttell process "Microsoft Excel"
\t\ttell window 1
\t\t\tset p to position
\t\t\tset s to size
\t\t\treturn ((item 1 of p) as text) & "," & ((item 2 of p) as text) & "|" & ((item 1 of s) as text) & "," & ((item 2 of s) as text)
\t\tend tell
\tend tell
end tell
"""


WINDOW_NAME_SCRIPT = """
tell application "Microsoft Excel" to activate
delay 0.3
tell application "System Events"
\ttell process "Microsoft Excel"
\t\tset windowNames to name of every window
\t\treturn item 1 of windowNames as text
\tend tell
end tell
"""


JXA_MOUSE_SCRIPT = """
ObjC.import('ApplicationServices');
ObjC.import('Cocoa');
ObjC.import('Foundation');

function currentPoint() {
  return $.NSEvent.mouseLocation;
}

function mainScreenHeight() {
  const screen = $.NSScreen.screens.objectAtIndex(0);
  return Number(screen.frame.size.height);
}

function toQuartzPoint(x, y) {
  return $.CGPointMake(x, mainScreenHeight() - y);
}

function sleepSeconds(seconds) {
  $.NSThread.sleepForTimeInterval(seconds);
}

function clickAt(x, y) {
  const point = toQuartzPoint(x, y);
  const moveEvent = $.CGEventCreateMouseEvent(null, $.kCGEventMouseMoved, point, $.kCGMouseButtonLeft);
  const downEvent = $.CGEventCreateMouseEvent(null, $.kCGEventLeftMouseDown, point, $.kCGMouseButtonLeft);
  const upEvent = $.CGEventCreateMouseEvent(null, $.kCGEventLeftMouseUp, point, $.kCGMouseButtonLeft);
  $.CGEventPost($.kCGHIDEventTap, moveEvent);
  sleepSeconds(0.15);
  $.CGEventPost($.kCGHIDEventTap, downEvent);
  sleepSeconds(0.10);
  $.CGEventPost($.kCGHIDEventTap, upEvent);
}

function run(argv) {
  const action = argv[0];
  if (action === "location") {
    const point = currentPoint();
    return `${point.x},${point.y}`;
  }
  if (action === "click") {
    const x = Number(argv[1]);
    const y = Number(argv[2]);
    clickAt(x, y);
    return `clicked:${x},${y}`;
  }
  throw new Error(`Unknown action: ${action}`);
}
"""


SUBMIT_QUERY_SCRIPT = """
on run argv
\tset queryText to item 1 of argv
\tset interDelay to (item 2 of argv) as real
\ttell application "Microsoft Excel" to activate
\tdelay 0.2
\ttell application "System Events"
\t\ttell process "Microsoft Excel"
\t\t\tset frontmost to true
\t\t\tkeystroke "a" using command down
\t\t\tdelay interDelay
\t\t\tkeystroke queryText
\t\tend tell
\tend tell
\treturn "ok"
end run
"""


FETCH_SNAPSHOT_SCRIPT = """
on run argv
\tset workbookName to item 1 of argv
\tset worksheetName to item 2 of argv
\tset maxRows to (item 3 of argv) as integer
\tset maxColumns to (item 4 of argv) as integer
\tset rowSep to character id 30
\tset cellSep to character id 31
\tset emptyMarker to character id 29
\t
\ttell application "Microsoft Excel"
\t\tif worksheetName is "active" then
\t\t\tset targetSheet to active sheet
\t\telse
\t\t\tset targetSheet to worksheet (worksheetName as string) of workbook (workbookName as string)
\t\tend if
\t\ttell targetSheet
\t\t\tset rowCount to maxRows
\t\t\tset columnCount to maxColumns
\t\t\tset lastColumnLetter to my columnLetter(columnCount)
\t\t\tset dataRows to value of range ("A1:" & lastColumnLetter & rowCount)
\t\tend tell
\tend tell
\t
\tset outputRows to {}
\trepeat with rowValues in dataRows
\t\tset cellTexts to {}
\t\trepeat with rawCell in rowValues
\t\t\tif rawCell is missing value then
\t\t\t\tset end of cellTexts to emptyMarker
\t\t\telse
\t\t\t\tset end of cellTexts to (rawCell as text)
\t\t\tend if
\t\tend repeat
\t\t
\t\tset oldTID to AppleScript's text item delimiters
\t\tset AppleScript's text item delimiters to cellSep
\t\tset end of outputRows to (cellTexts as text)
\t\tset AppleScript's text item delimiters to oldTID
\tend repeat
\t
\tset oldTID to AppleScript's text item delimiters
\tset AppleScript's text item delimiters to rowSep
\tset outputText to outputRows as text
\tset AppleScript's text item delimiters to oldTID
\t
\treturn ((rowCount as text) & rowSep & (columnCount as text) & rowSep & outputText)
end run

on columnLetter(columnNumber)
\tset dividend to columnNumber
\tset columnName to ""
\trepeat while dividend > 0
\t\tset remainder to (dividend - 1) mod 26
\t\tset columnName to (character (remainder + 1) of "ABCDEFGHIJKLMNOPQRSTUVWXYZ") & columnName
\t\tset dividend to (dividend - remainder - 1) div 26
\tend repeat
\treturn columnName
end columnLetter
"""


FETCH_RANGE_CHUNK_SCRIPT = """
on run argv
\tset workbookName to item 1 of argv
\tset worksheetName to item 2 of argv
\tset startRow to (item 3 of argv) as integer
\tset endRow to (item 4 of argv) as integer
\tset columnCount to (item 5 of argv) as integer
\tset rowSep to character id 30
\tset cellSep to character id 31
\tset emptyMarker to character id 29
\t
\ttell application "Microsoft Excel"
\t\tif worksheetName is "active" then
\t\t\tset targetSheet to active sheet
\t\telse
\t\t\tset targetSheet to worksheet (worksheetName as string) of workbook (workbookName as string)
\t\tend if
\t\ttell targetSheet
\t\t\tset lastColumnLetter to my columnLetter(columnCount)
\t\t\tset dataRows to value of range ("A" & startRow & ":" & lastColumnLetter & endRow)
\t\tend tell
\tend tell
\t
\tset outputRows to {}
\trepeat with rowValues in dataRows
\t\tset cellTexts to {}
\t\trepeat with rawCell in rowValues
\t\t\tif rawCell is missing value then
\t\t\t\tset end of cellTexts to emptyMarker
\t\t\telse
\t\t\t\tset end of cellTexts to (rawCell as text)
\t\t\tend if
\t\tend repeat
\t\tset oldTID to AppleScript's text item delimiters
\t\tset AppleScript's text item delimiters to cellSep
\t\tset end of outputRows to (cellTexts as text)
\t\tset AppleScript's text item delimiters to oldTID
\tend repeat
\t
\tset oldTID to AppleScript's text item delimiters
\tset AppleScript's text item delimiters to rowSep
\tset outputText to outputRows as text
\tset AppleScript's text item delimiters to oldTID
\treturn outputText
end run

on columnLetter(columnNumber)
\tset dividend to columnNumber
\tset columnName to ""
\trepeat while dividend > 0
\t\tset remainder to (dividend - 1) mod 26
\t\tset columnName to (character (remainder + 1) of "ABCDEFGHIJKLMNOPQRSTUVWXYZ") & columnName
\t\tset dividend to (dividend - remainder - 1) div 26
\tend repeat
\treturn columnName
end columnLetter
"""


FETCH_SHEET_NAME_SCRIPT = """
on run argv
\tset workbookName to item 1 of argv
\tset worksheetName to item 2 of argv
\ttell application "Microsoft Excel"
\t\treturn name of worksheet (worksheetName as string) of workbook (workbookName as string)
\tend tell
end run
"""


DUPLICATE_SHEET_TO_WORKBOOK_SCRIPT = """
on run argv
\tset sourceWorkbookName to item 1 of argv
\tset sourceWorksheetName to item 2 of argv
\tset outputWorkbookPath to item 3 of argv
\tset outputWorkbookName to item 4 of argv
\tset outputSheetBaseName to item 5 of argv
\tset outputWorkbookFile to POSIX file outputWorkbookPath
\t
\ttell application "Microsoft Excel"
\t\tactivate
\t\tif not (exists workbook (outputWorkbookName as string)) then
\t\t\topen outputWorkbookFile
\t\t\tdelay 0.4
\t\tend if
\t\t
\t\tset sourceSheet to worksheet (sourceWorksheetName as string) of workbook (sourceWorkbookName as string)
\t\tset destinationBook to workbook (outputWorkbookName as string)
\t\tset destinationSheetCountBefore to count of worksheets of destinationBook
\t\t
\t\tcopy sourceSheet to destinationBook
\t\tdelay 0.3
\t\tset copiedSheet to worksheet (destinationSheetCountBefore + 1) of destinationBook
\t\tset existingNames to name of worksheets of destinationBook
\t\tset copiedSheetName to name of copiedSheet
\t\tset outputSheetName to my uniqueWorksheetName(outputSheetBaseName, existingNames, copiedSheetName)
\t\tset name of copiedSheet to outputSheetName
\t\tif (count of worksheets of destinationBook) > 1 then
\t\t\ttry
\t\t\t\tdelete worksheet ("_placeholder" as string) of destinationBook
\t\t\tend try
\t\tend if
\t\tclose destinationBook saving yes
\tend tell
\t
\treturn outputSheetName
end run

on uniqueWorksheetName(baseName, existingNames, copiedSheetName)
\tif copiedSheetName is in existingNames then
\t\tset existingNames to my removeItem(existingNames, copiedSheetName)
\tend if
\t
\tset candidateName to my truncateSheetName(baseName)
\tif candidateName is not in existingNames then
\t\treturn candidateName
\tend if
\t
\tset suffix to 2
\trepeat
\t\tset suffixText to "_" & (suffix as text)
\t\tset baseLimit to 31 - (length of suffixText)
\t\tset candidatePrefix to my truncateSheetName(baseName)
\t\tif (length of candidatePrefix) > baseLimit then
\t\t\tset candidatePrefix to text 1 thru baseLimit of candidatePrefix
\t\tend if
\t\tset candidateName to candidatePrefix & suffixText
\t\tif candidateName is not in existingNames then
\t\t\treturn candidateName
\t\tend if
\t\tset suffix to suffix + 1
\tend repeat
end uniqueWorksheetName

on truncateSheetName(baseName)
\tif (length of baseName) > 31 then
\t\treturn text 1 thru 31 of baseName
\tend if
\treturn baseName
end truncateSheetName

on removeItem(theList, itemToRemove)
\tset newList to {}
\trepeat with currentItem in theList
\t\tif (currentItem as text) is not (itemToRemove as text) then
\t\t\tset end of newList to (currentItem as text)
\t\tend if
\tend repeat
\treturn newList
end removeItem
"""


ACTIVATE_WORKBOOK_SHEET_SCRIPT = """
on run argv
\tset workbookName to item 1 of argv
\tset worksheetName to item 2 of argv
\ttell application "Microsoft Excel"
\t\tactivate
\t\tactivate object worksheet (worksheetName as string) of workbook (workbookName as string)
\tend tell
\treturn "ok"
end run
"""


@dataclass(frozen=True)
class QueryItem:
    source_row: int
    query_text: str
    company_name: str
    inn: str


@dataclass(frozen=True)
class WindowGeometry:
    name: str
    x: float
    y: float
    width: float
    height: float


@dataclass(frozen=True)
class RelativePoint:
    x_ratio: float
    y_ratio: float

    def to_absolute(self, geometry: WindowGeometry) -> tuple[int, int]:
        return (
            round(geometry.x + geometry.width * self.x_ratio),
            round(geometry.y + geometry.height * self.y_ratio),
        )


REQUIRED_UI_POINTS = ("search_input", "first_result", "submit_button")


def run_osascript(script: str, *args: str) -> str:
    result = subprocess.run(
        ["osascript", "-", *args],
        input=script,
        text=True,
        capture_output=True,
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or f"osascript failed: {result.returncode}")
    return result.stdout.rstrip("\n")


def run_jxa(script: str, *args: str) -> str:
    result = subprocess.run(
        ["osascript", "-l", "JavaScript", "-", *args],
        input=script,
        text=True,
        capture_output=True,
        check=False,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or f"JXA failed: {result.returncode}")
    return result.stdout.rstrip("\n")


def clean_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def ensure_accessibility_enabled() -> None:
    raw = run_osascript(ACCESSIBILITY_CHECK_SCRIPT).strip().lower()
    if raw != "true":
        raise RuntimeError("Accessibility не включен для Terminal/osascript.")


def ensure_cbonds_sheet_open(workbook_name: str, worksheet_name: str) -> None:
    run_osascript(WORKBOOK_SHEET_CHECK_SCRIPT, workbook_name, worksheet_name)


def list_worksheets(workbook_name: str) -> list[str]:
    raw = run_osascript(LIST_WORKSHEETS_SCRIPT, workbook_name).strip()
    return [line.strip() for line in raw.splitlines() if line.strip()]


def get_excel_window_geometry() -> WindowGeometry:
    raw_name = run_osascript(WINDOW_NAME_SCRIPT).strip()
    raw_geometry = run_osascript(WINDOW_GEOMETRY_SCRIPT).strip()
    position_text, size_text = raw_geometry.split("|")
    x_text, y_text = position_text.split(",")
    width_text, height_text = size_text.split(",")
    return WindowGeometry(
        name=raw_name,
        x=float(x_text),
        y=float(y_text),
        width=float(width_text),
        height=float(height_text),
    )


def ensure_expected_excel_window(expected_name: str) -> WindowGeometry:
    geometry = get_excel_window_geometry()
    if geometry.name and geometry.name != expected_name:
        raise RuntimeError(
            f"Переднее окно Excel сейчас {geometry.name!r}, ожидалось {expected_name!r}. "
            "Выведите нужную книгу на передний план."
        )
    return geometry


def get_mouse_location() -> tuple[float, float]:
    raw = run_jxa(JXA_MOUSE_SCRIPT, "location").strip()
    x_text, y_text = raw.split(",")
    return float(x_text), float(y_text)


def click_relative_point(geometry: WindowGeometry, point: RelativePoint) -> tuple[int, int]:
    x, y = point.to_absolute(geometry)
    run_jxa(JXA_MOUSE_SCRIPT, "click", str(x), str(y))
    return x, y


def column_letter(column_number: int) -> str:
    column_name = ""
    dividend = column_number
    while dividend > 0:
        dividend, remainder = divmod(dividend - 1, 26)
        column_name = chr(65 + remainder) + column_name
    return column_name


def load_queries(source_csv: Path, query_header: str, limit: int | None, offset: int) -> list[QueryItem]:
    items: list[QueryItem] = []
    skipped = 0
    with source_csv.open("r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.DictReader(fh)
        if not reader.fieldnames or query_header not in reader.fieldnames:
            raise RuntimeError(f"Не нашел колонку {query_header!r} в {source_csv}")
        for row_number, row in enumerate(reader, start=2):
            query = clean_text(row.get(query_header))
            if not query:
                continue
            if skipped < offset:
                skipped += 1
                continue
            items.append(
                QueryItem(
                    source_row=row_number,
                    query_text=query,
                    company_name=clean_text(row.get("company_name")),
                    inn=clean_text(row.get("inn")),
                )
            )
            if limit is not None and len(items) >= limit:
                break
    return items


def load_ui_points(path: Path) -> dict[str, RelativePoint]:
    if not path.exists():
        raise RuntimeError(f"Не найден файл калибровки UI: {path}. Сначала запустите --calibrate-ui.")
    payload = json.loads(path.read_text(encoding="utf-8"))
    points_payload = payload.get("points", {})
    missing = [key for key in REQUIRED_UI_POINTS if key not in points_payload]
    if missing:
        raise RuntimeError(f"В калибровке не хватает точек: {', '.join(missing)}")
    return {
        key: RelativePoint(
            x_ratio=float(points_payload[key]["x_ratio"]),
            y_ratio=float(points_payload[key]["y_ratio"]),
        )
        for key in REQUIRED_UI_POINTS
    }


def save_ui_points(path: Path, geometry: WindowGeometry, points: dict[str, RelativePoint]) -> None:
    payload = {
        "calibrated_at": datetime.now().isoformat(timespec="seconds"),
        "window_name": geometry.name,
        "window_size": {"width": geometry.width, "height": geometry.height},
        "points": {key: {"x_ratio": point.x_ratio, "y_ratio": point.y_ratio} for key, point in points.items()},
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def capture_relative_point(label: str, workbook_name: str, countdown_seconds: int) -> RelativePoint:
    input(
        f"{label}\n"
        f"Нажмите Enter, затем Excel выйдет на передний план, и у вас будет {countdown_seconds} сек., "
        "чтобы навести курсор в центр нужного элемента..."
    )
    geometry = ensure_expected_excel_window(workbook_name)
    for seconds_left in range(countdown_seconds, 0, -1):
        print(f"  capture in {seconds_left}...", flush=True)
        time.sleep(1)
    mouse_x, mouse_y = get_mouse_location()
    x_ratio = (mouse_x - geometry.x) / geometry.width
    y_ratio = (mouse_y - geometry.y) / geometry.height
    if not (0 <= x_ratio <= 1 and 0 <= y_ratio <= 1):
        raise RuntimeError("Курсор оказался вне окна Excel во время калибровки.")
    print(f"  captured absolute=({mouse_x:.1f}, {mouse_y:.1f}) relative=({x_ratio:.4f}, {y_ratio:.4f})", flush=True)
    return RelativePoint(x_ratio=x_ratio, y_ratio=y_ratio)


def calibrate_ui_points(args: argparse.Namespace) -> None:
    print("Калибровка UI для карточек компаний Cbonds")
    points = {
        "search_input": capture_relative_point(
            "1/3. Наведите курсор на центр поля поиска компаний.",
            args.cbonds_workbook,
            args.calibration_countdown,
        ),
        "first_result": capture_relative_point(
            "2/3. Введите любой запрос вручную, дождитесь dropdown и наведите курсор на первую найденную компанию.",
            args.cbonds_workbook,
            args.calibration_countdown,
        ),
        "submit_button": capture_relative_point(
            "3/3. Наведите курсор на центр кнопки 'Получить данные'.",
            args.cbonds_workbook,
            args.calibration_countdown,
        ),
    }
    save_ui_points(args.ui_config, ensure_expected_excel_window(args.cbonds_workbook), points)
    print(f"Калибровка сохранена: {args.ui_config}")


def submit_query_via_click_workflow(
    query_text: str,
    *,
    workbook_name: str,
    ui_points: dict[str, RelativePoint],
    inter_key_delay: float,
    click_delay: float,
    suggestion_wait_seconds: float,
    after_result_wait_seconds: float,
) -> None:
    print("  -> focusing search input", flush=True)
    geometry = ensure_expected_excel_window(workbook_name)
    click_relative_point(geometry, ui_points["search_input"])
    time.sleep(click_delay)

    print("  -> typing query", flush=True)
    run_osascript(SUBMIT_QUERY_SCRIPT, query_text, str(inter_key_delay))
    print(f"  -> waiting {suggestion_wait_seconds:.1f}s for dropdown", flush=True)
    time.sleep(suggestion_wait_seconds)

    print("  -> clicking first result", flush=True)
    click_relative_point(geometry, ui_points["first_result"])
    print(f"  -> waiting {after_result_wait_seconds:.1f}s after result click", flush=True)
    time.sleep(after_result_wait_seconds)

    print("  -> clicking submit button", flush=True)
    geometry = ensure_expected_excel_window(workbook_name)
    click_relative_point(geometry, ui_points["submit_button"])


def fetch_sheet_matrix(
    workbook_name: str,
    worksheet_name: str,
    max_rows: int,
    max_columns: int,
    chunk_rows: int,
) -> list[list[str | None]]:
    matrix: list[list[str | None]] = []
    for start_row in range(1, max_rows + 1, chunk_rows):
        end_row = min(start_row + chunk_rows - 1, max_rows)
        raw = run_osascript(
            FETCH_RANGE_CHUNK_SCRIPT,
            workbook_name,
            worksheet_name,
            str(start_row),
            str(end_row),
            str(max_columns),
        )
        for row_text in raw.split(ROW_SEPARATOR):
            row_values = row_text.split(CELL_SEPARATOR) if row_text else []
            normalized = [None if cell == EMPTY_MARKER or cell == "" else cell for cell in row_values[:max_columns]]
            normalized.extend([None] * (max_columns - len(normalized)))
            matrix.append(normalized)
    return matrix


def sanitize_sheet_title(text: str) -> str:
    cleaned = re.sub(r'[:\\\\/?*\\[\\]]+', "_", text).strip()
    cleaned = re.sub(r"_+", "_", cleaned)
    return cleaned or "company"


def build_unique_sheet_title(wb: Workbook, item: QueryItem) -> str:
    base = f"r{item.source_row}_{sanitize_sheet_title(item.query_text)}"
    base = base[:31]
    if base not in wb.sheetnames:
        return base
    suffix = 2
    while True:
        suffix_text = f"_{suffix}"
        candidate = base[: 31 - len(suffix_text)] + suffix_text
        if candidate not in wb.sheetnames:
            return candidate
        suffix += 1


def build_base_sheet_title(item: QueryItem) -> str:
    return f"r{item.source_row}_{sanitize_sheet_title(item.query_text)}"[:31]


def load_or_create_output_workbook(path: Path) -> Workbook:
    if path.exists():
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.active.title = "capture_log"
        wb["capture_log"].append(["timestamp", "source_row", "query_text", "company_name", "inn", "sheet_name", "status"])
    if "capture_log" not in wb.sheetnames:
        ws = wb.create_sheet("capture_log", 0)
        ws.append(["timestamp", "source_row", "query_text", "company_name", "inn", "sheet_name", "status"])
    return wb


def ensure_duplicate_output_workbook(path: Path) -> None:
    if path.exists():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.active.title = "_placeholder"
    wb.save(path)


def duplicate_sheet_to_output_workbook(
    *,
    source_workbook_name: str,
    source_sheet_name: str,
    output_workbook: Path,
    output_sheet_name: str,
) -> str:
    ensure_duplicate_output_workbook(output_workbook)
    last_error: RuntimeError | None = None
    for attempt in range(1, 6):
        try:
            return run_osascript(
                DUPLICATE_SHEET_TO_WORKBOOK_SCRIPT,
                source_workbook_name,
                source_sheet_name,
                str(output_workbook),
                output_workbook.name,
                output_sheet_name,
            )
        except RuntimeError as exc:
            last_error = exc
            if attempt == 5:
                break
            print(f"  -> duplicate attempt {attempt} failed, retrying in 1s: {exc}", flush=True)
            time.sleep(1)
    if last_error is not None:
        raise last_error
    raise RuntimeError("Не удалось продублировать лист Cbonds.")


def activate_workbook_sheet(workbook_name: str, worksheet_name: str) -> None:
    run_osascript(ACTIVATE_WORKBOOK_SHEET_SCRIPT, workbook_name, worksheet_name)


def write_snapshot_sheet(wb: Workbook, item: QueryItem, matrix: list[list[str | None]]) -> str:
    sheet_name = build_unique_sheet_title(wb, item)
    ws = wb.create_sheet(sheet_name)
    for row_index, row_values in enumerate(matrix, start=1):
        for column_index, value in enumerate(row_values, start=1):
            if value is not None:
                ws.cell(row=row_index, column=column_index, value=value)
    return sheet_name


def append_log_entry(wb: Workbook, item: QueryItem, sheet_name: str, status: str) -> None:
    wb["capture_log"].append([
        datetime.now().isoformat(timespec="seconds"),
        item.source_row,
        item.query_text,
        item.company_name,
        item.inn,
        sheet_name,
        status,
    ])


def prompt_before_start(items: list[QueryItem], args: argparse.Namespace) -> None:
    print("Подготовка к прогону Cbonds company cards")
    print(f"Источник: {args.source_csv}")
    print(f"Колонка запроса: {args.query_header}")
    print(f"Книга Cbonds: {args.cbonds_workbook} / {args.cbonds_sheet}")
    print(f"Выходной файл: {args.output_workbook}")
    print()
    print("Первые запросы:")
    for item in items[: min(len(items), 5)]:
        print(f"  row {item.source_row}: {item.query_text} | {item.company_name}")
    print()
    print("Перед стартом:")
    print("1. Откройте Excel и книгу с Cbonds.")
    print("2. Выведите книгу Cbonds на передний план.")
    print("3. Не меняйте размер окна Excel во время прогона.")
    print("4. Не переключайтесь из Excel, пока идет тест.")
    input("Нажмите Enter в терминале, когда будете готовы...")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Capture Cbonds company cards from the Office add-in on Mac.")
    parser.add_argument("--source-csv", type=Path, default=DEFAULT_SOURCE_CSV)
    parser.add_argument("--query-header", default="inn")
    parser.add_argument("--cbonds-workbook", default="Книга2")
    parser.add_argument(
        "--cbonds-sheet",
        default="CbondsCalendar",
        help="Имя листа Cbonds для чтения.",
    )
    parser.add_argument("--output-workbook", type=Path, default=DEFAULT_OUTPUT_WORKBOOK)
    parser.add_argument("--ui-config", type=Path, default=DEFAULT_UI_CONFIG)
    parser.add_argument("--offset", type=int, default=28)
    parser.add_argument("--limit", type=int, default=1)
    parser.add_argument("--wait-seconds", type=float, default=2.0)
    parser.add_argument(
        "--capture-mode",
        choices=["duplicate_sheet", "matrix"],
        default="duplicate_sheet",
        help="Как сохранять результат: дублировать весь лист Excel или читать матрицу значений.",
    )
    parser.add_argument("--max-copy-rows", type=int, default=2000)
    parser.add_argument("--max-copy-columns", type=int, default=100)
    parser.add_argument("--copy-chunk-rows", type=int, default=50)
    parser.add_argument("--inter-key-delay", type=float, default=0.15)
    parser.add_argument("--click-delay", type=float, default=0.20)
    parser.add_argument("--suggestion-wait-seconds", type=float, default=2.5)
    parser.add_argument("--after-result-wait-seconds", type=float, default=0.80)
    parser.add_argument("--calibration-countdown", type=int, default=3)
    parser.add_argument("--calibrate-ui", action="store_true")
    parser.add_argument("--list-sheets", action="store_true")
    parser.add_argument("--skip-submit", action="store_true")
    parser.add_argument("--yes", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    print("Проверяю Accessibility...", flush=True)
    ensure_accessibility_enabled()
    print("Проверяю переднее окно Excel...", flush=True)
    ensure_expected_excel_window(args.cbonds_workbook)

    if args.list_sheets:
        print(f"Листы в книге {args.cbonds_workbook}:")
        for worksheet_name in list_worksheets(args.cbonds_workbook):
            print(f"  - {worksheet_name}")
        return

    if args.calibrate_ui:
        calibrate_ui_points(args)
        return

    if args.cbonds_sheet != "active":
        print("Проверяю лист Cbonds...", flush=True)
        try:
            ensure_cbonds_sheet_open(args.cbonds_workbook, args.cbonds_sheet)
        except RuntimeError as exc:
            worksheet_hint = ""
            try:
                worksheet_names = list_worksheets(args.cbonds_workbook)
                if worksheet_names:
                    worksheet_hint = "\nДоступные листы:\n" + "\n".join(f"  - {name}" for name in worksheet_names)
            except RuntimeError:
                worksheet_hint = ""
            raise RuntimeError(
                f"Не нашел лист {args.cbonds_sheet!r} в книге {args.cbonds_workbook!r}."
                f"{worksheet_hint}\nПередайте точное имя листа через --cbonds-sheet или оставьте active."
            ) from exc

    items = load_queries(args.source_csv, args.query_header, args.limit, args.offset)
    if not items:
        raise RuntimeError("Не нашлось непустых запросов для прогона.")

    ui_points = load_ui_points(args.ui_config)
    if not args.yes:
        prompt_before_start(items, args)

    args.output_workbook.parent.mkdir(parents=True, exist_ok=True)
    output_wb = None
    if args.capture_mode == "matrix":
        output_wb = load_or_create_output_workbook(args.output_workbook)
    else:
        ensure_duplicate_output_workbook(args.output_workbook)

    for index, item in enumerate(items, start=1):
        print(f"[{index}/{len(items)}] {item.query_text} | {item.company_name}", flush=True)
        if not args.skip_submit:
            submit_query_via_click_workflow(
                item.query_text,
                workbook_name=args.cbonds_workbook,
                ui_points=ui_points,
                inter_key_delay=args.inter_key_delay,
                click_delay=args.click_delay,
                suggestion_wait_seconds=args.suggestion_wait_seconds,
                after_result_wait_seconds=args.after_result_wait_seconds,
            )
            print(f"  -> waiting {args.wait_seconds:.1f}s for sheet refresh", flush=True)
            time.sleep(args.wait_seconds)

        sheet_name = build_base_sheet_title(item)
        if args.capture_mode == "duplicate_sheet":
            print(f"  -> duplicating sheet {args.cbonds_sheet} into {args.output_workbook.name}", flush=True)
            sheet_name = duplicate_sheet_to_output_workbook(
                source_workbook_name=args.cbonds_workbook,
                source_sheet_name=args.cbonds_sheet,
                output_workbook=args.output_workbook,
                output_sheet_name=sheet_name,
            )
            print(f"  -> returning focus to {args.cbonds_workbook}/{args.cbonds_sheet}", flush=True)
            try:
                activate_workbook_sheet(args.cbonds_workbook, args.cbonds_sheet)
            except RuntimeError as exc:
                print(f"  -> focus return skipped: {exc}", flush=True)
        else:
            if output_wb is None:
                raise RuntimeError("output_wb is not initialized for matrix capture mode.")
            print(f"  -> copying sheet {args.cbonds_sheet} up to {args.max_copy_rows}x{args.max_copy_columns}", flush=True)
            matrix = fetch_sheet_matrix(
                args.cbonds_workbook,
                args.cbonds_sheet,
                max_rows=args.max_copy_rows,
                max_columns=args.max_copy_columns,
                chunk_rows=args.copy_chunk_rows,
            )
            sheet_name = write_snapshot_sheet(output_wb, item, matrix)
            append_log_entry(output_wb, item, sheet_name, "captured")
            output_wb.save(args.output_workbook)
        print(f"  -> saved to sheet {sheet_name}", flush=True)

    print(f"Готово: {args.output_workbook}")


if __name__ == "__main__":
    main()
