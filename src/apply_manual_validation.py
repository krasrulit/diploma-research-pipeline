from __future__ import annotations

import argparse
import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font


MAIN_WORKBOOK_DEFAULT = Path("data_processed/additional_public_data.xlsx")
VALIDATION_WORKBOOK_DEFAULT = Path("data_processed/additional_public_data_manual_review_validated.xlsx")
OUTPUT_WORKBOOK_DEFAULT = Path("data_processed/additional_public_data_validated.xlsx")


ISIN_RE = re.compile(r"\b[A-Z]{2}[A-Z0-9]{10}\b")


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    if text.lower() == "nan":
        return ""
    return text


def bool_value(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return clean_text(value).lower() in {"true", "1", "yes", "да"}


def combine_text(row: pd.Series) -> str:
    fields = [
        "issue_summary",
        "current_problem",
        "what_is_correct",
        "what_is_wrong",
        "what_should_be_changed",
        "notes",
        "source_urls",
    ]
    return " ".join(clean_text(row.get(field, "")) for field in fields)


def extract_isins(text: str) -> list[str]:
    return sorted(set(ISIN_RE.findall(text.upper())))


def extract_known_ids(text: str, known_ids: set[str]) -> list[str]:
    upper_text = text.upper()
    found = []
    for instrument_id in known_ids:
        if not instrument_id:
            continue
        pattern = rf"(?<![A-Z0-9]){re.escape(instrument_id)}(?![A-Z0-9])"
        if re.search(pattern, upper_text):
            found.append(instrument_id)
    return sorted(set(found))


def classify_tinvest_action(row: pd.Series) -> str:
    verdict = clean_text(row.get("verdict", "")).upper()
    text = combine_text(row).lower()

    if verdict == "REJECT":
        return "reject_tinvest_candidates"
    if verdict == "UNCERTAIN":
        return "do_not_use_until_resolved"
    if verdict == "NO_TINVEST_COVERAGE":
        return "use_moex_coverage_gap"
    if verdict == "CONFIRM":
        false_candidate_markers = [
            "t-invest matched",
            "t‑invest matched",
            "unrelated",
            "disregard",
            "ignore false",
            "false candidate",
            "false matches",
            "map company to moex",
        ]
        if any(marker in text for marker in false_candidate_markers):
            return "use_moex_mapping_not_tinvest"
        return "confirm_tinvest_mapping"
    return "manual_review"


def load_validation(validation_workbook: Path) -> pd.DataFrame:
    frame = pd.read_excel(validation_workbook, sheet_name="validated_checks")
    required = ["company_name", "review_bucket", "verdict", "confidence"]
    missing = [column for column in required if column not in frame.columns]
    if missing:
        raise ValueError(f"Missing required validation columns: {missing}")

    frame = frame.copy()
    frame["company_name"] = frame["company_name"].map(clean_text)
    frame["review_bucket"] = frame["review_bucket"].map(clean_text)
    frame["verdict"] = frame["verdict"].map(lambda value: clean_text(value).upper())
    frame["confidence"] = frame["confidence"].map(clean_text)
    frame["validation_action"] = frame.apply(classify_tinvest_action, axis=1)
    frame["validation_text"] = frame.apply(combine_text, axis=1)
    return frame


def build_known_ids(
    moex_instruments: pd.DataFrame,
    moex_bonds: pd.DataFrame,
    tinvest_instruments: pd.DataFrame,
) -> set[str]:
    values: set[str] = set()
    for frame, columns in [
        (moex_instruments, ["secid", "isin"]),
        (moex_bonds, ["SECID", "ISIN"]),
        (tinvest_instruments, ["ticker", "isin"]),
    ]:
        for column in columns:
            if column not in frame.columns:
                continue
            for value in frame[column].dropna().astype(str):
                instrument_id = clean_text(value).upper()
                if len(instrument_id) >= 3:
                    values.add(instrument_id)
    return values


def build_validated_actions(validation_df: pd.DataFrame, known_ids: set[str]) -> pd.DataFrame:
    actions = validation_df.copy()
    actions["validated_isins"] = actions["validation_text"].map(lambda text: ", ".join(extract_isins(text)))
    actions["validated_ids"] = actions["validation_text"].map(lambda text: ", ".join(extract_known_ids(text, known_ids)))
    actions["manual_use_tinvest"] = actions["validation_action"].eq("confirm_tinvest_mapping")
    actions["manual_use_moex"] = actions["validation_action"].isin(
        ["use_moex_mapping_not_tinvest", "use_moex_coverage_gap"]
    )
    actions["manual_reject"] = actions["validation_action"].eq("reject_tinvest_candidates")
    actions["manual_uncertain"] = actions["validation_action"].eq("do_not_use_until_resolved")
    return actions


def get_tinvest_review_actions(actions: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "company_name",
        "verdict",
        "confidence",
        "validation_action",
        "what_is_correct",
        "what_is_wrong",
        "what_should_be_changed",
        "notes",
        "manual_use_tinvest",
        "manual_use_moex",
        "manual_reject",
        "manual_uncertain",
    ]
    frame = actions[actions["review_bucket"].eq("tinvest_review")].copy()
    frame = frame[[column for column in columns if column in frame.columns]]
    return frame.drop_duplicates(subset=["company_name"], keep="first")


def annotate_tinvest_instruments(tinvest_instruments: pd.DataFrame, actions: pd.DataFrame) -> pd.DataFrame:
    if tinvest_instruments.empty:
        return tinvest_instruments.copy()

    action_cols = get_tinvest_review_actions(actions)
    annotated = tinvest_instruments.merge(
        action_cols,
        on="company_name",
        how="left",
        suffixes=("", "_manual"),
    )
    annotated = annotated.rename(
        columns={
            "verdict": "manual_verdict",
            "confidence": "manual_confidence",
            "what_is_correct": "manual_what_is_correct",
            "what_is_wrong": "manual_what_is_wrong",
            "what_should_be_changed": "manual_what_should_be_changed",
            "notes": "manual_notes",
        }
    )
    annotated["manual_tinvest_keep_flag"] = annotated.apply(
        lambda row: bool_value(row.get("selected_flag", "")) if bool_value(row.get("manual_use_tinvest", "")) else False,
        axis=1,
    )
    annotated["manual_review_needed_flag"] = annotated["manual_verdict"].fillna("").astype(str).ne("")
    return annotated


def validation_key_rows(actions: pd.DataFrame, action_name: str) -> pd.DataFrame:
    return actions[actions["validation_action"].eq(action_name)].copy()


def split_ids(value: object) -> set[str]:
    return {item.strip().upper() for item in clean_text(value).split(",") if item.strip()}


def rows_for_ids(frame: pd.DataFrame, company_name: str, id_columns: list[str], wanted_ids: set[str]) -> pd.DataFrame:
    if frame.empty or not wanted_ids:
        return pd.DataFrame()
    subset = frame[frame["company_name"].eq(company_name)].copy()
    if subset.empty:
        return pd.DataFrame()

    mask = pd.Series(False, index=subset.index)
    for column in id_columns:
        if column not in subset.columns:
            continue
        mask = mask | subset[column].astype(str).str.upper().isin(wanted_ids)
    return subset[mask].copy()


def normalize_mapping_row(
    row: pd.Series,
    instrument_source: str,
    company_name: str,
    validation_row: pd.Series,
) -> dict[str, object]:
    secid = clean_text(row.get("secid", row.get("SECID", row.get("ticker", ""))))
    isin = clean_text(row.get("isin", row.get("ISIN", "")))
    instrument_name = clean_text(row.get("name", row.get("SECNAME", row.get("SHORTNAME", ""))))
    instrument_kind = clean_text(row.get("type", row.get("instrument_kind_source", row.get("GROUP", ""))))

    return {
        "company_name": company_name,
        "instrument_source": instrument_source,
        "instrument_id": secid,
        "isin": isin,
        "instrument_name": instrument_name,
        "instrument_kind": instrument_kind,
        "validation_verdict": clean_text(validation_row.get("verdict", "")),
        "validation_confidence": clean_text(validation_row.get("confidence", "")),
        "validation_action": clean_text(validation_row.get("validation_action", "")),
        "review_bucket": clean_text(validation_row.get("review_bucket", "")),
        "what_should_be_changed": clean_text(validation_row.get("what_should_be_changed", "")),
        "notes": clean_text(validation_row.get("notes", "")),
    }


def placeholder_mapping_rows(validation_row: pd.Series, company_name: str, ids: set[str], source: str) -> list[dict[str, object]]:
    rows = []
    for instrument_id in sorted(ids):
        rows.append(
            {
                "company_name": company_name,
                "instrument_source": source,
                "instrument_id": instrument_id,
                "isin": instrument_id if ISIN_RE.fullmatch(instrument_id) else "",
                "instrument_name": "",
                "instrument_kind": "",
                "validation_verdict": clean_text(validation_row.get("verdict", "")),
                "validation_confidence": clean_text(validation_row.get("confidence", "")),
                "validation_action": clean_text(validation_row.get("validation_action", "")),
                "review_bucket": clean_text(validation_row.get("review_bucket", "")),
                "what_should_be_changed": clean_text(validation_row.get("what_should_be_changed", "")),
                "notes": clean_text(validation_row.get("notes", "")),
            }
        )
    return rows


def build_validated_mapping(
    actions: pd.DataFrame,
    tinvest_instruments: pd.DataFrame,
    moex_instruments: pd.DataFrame,
    moex_bonds: pd.DataFrame,
) -> pd.DataFrame:
    rows: list[dict[str, object]] = []

    for _, action in actions.iterrows():
        company_name = clean_text(action.get("company_name", ""))
        action_name = clean_text(action.get("validation_action", ""))
        ids = split_ids(action.get("validated_ids", ""))
        isins = split_ids(action.get("validated_isins", ""))
        wanted_ids = ids | isins

        if action_name == "confirm_tinvest_mapping":
            subset = tinvest_instruments[
                tinvest_instruments["company_name"].eq(company_name)
                & tinvest_instruments.get("selected_flag", pd.Series(False, index=tinvest_instruments.index)).map(bool_value)
            ].copy()
            for _, row in subset.iterrows():
                rows.append(normalize_mapping_row(row, "T-Invest", company_name, action))

        elif action_name == "use_moex_mapping_not_tinvest":
            subset = rows_for_ids(moex_instruments, company_name, ["secid", "isin"], wanted_ids)
            if subset.empty:
                subset = moex_instruments[
                    moex_instruments["company_name"].eq(company_name)
                    & moex_instruments.get("selected_flag", pd.Series(False, index=moex_instruments.index)).map(bool_value)
                ].copy()
            for _, row in subset.iterrows():
                rows.append(normalize_mapping_row(row, "MOEX", company_name, action))

        elif action_name == "use_moex_coverage_gap":
            subset = rows_for_ids(moex_bonds, company_name, ["SECID", "ISIN"], wanted_ids)
            if subset.empty:
                subset = rows_for_ids(moex_instruments, company_name, ["secid", "isin"], wanted_ids)
            if subset.empty and wanted_ids:
                rows.extend(placeholder_mapping_rows(action, company_name, wanted_ids, "MOEX"))
            else:
                for _, row in subset.iterrows():
                    rows.append(normalize_mapping_row(row, "MOEX", company_name, action))

    result = pd.DataFrame(rows)
    if result.empty:
        return result
    return result.drop_duplicates(
        subset=["company_name", "instrument_source", "instrument_id", "isin", "validation_action"],
        keep="first",
    ).reset_index(drop=True)


def add_validated_flag(frame: pd.DataFrame, mapping: pd.DataFrame, id_columns: list[str]) -> pd.DataFrame:
    annotated = frame.copy()
    if annotated.empty or mapping.empty:
        annotated["manual_validated_mapping_flag"] = False
        return annotated

    mapping_keys = {
        (clean_text(row["company_name"]), clean_text(row["instrument_id"]).upper(), clean_text(row["isin"]).upper())
        for _, row in mapping.iterrows()
    }

    def is_validated(row: pd.Series) -> bool:
        company_name = clean_text(row.get("company_name", ""))
        ids = [clean_text(row.get(column, "")).upper() for column in id_columns]
        isin = clean_text(row.get("isin", row.get("ISIN", ""))).upper()
        for instrument_id in ids:
            if (company_name, instrument_id, isin) in mapping_keys:
                return True
            if any(key[0] == company_name and instrument_id and key[1] == instrument_id for key in mapping_keys):
                return True
            if any(key[0] == company_name and isin and key[2] == isin for key in mapping_keys):
                return True
        return False

    annotated["manual_validated_mapping_flag"] = annotated.apply(is_validated, axis=1)
    return annotated


def build_validation_summary(actions: pd.DataFrame, validated_mapping: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for column, label in [
        ("verdict", "Rows by verdict"),
        ("validation_action", "Rows by validation action"),
        ("review_bucket", "Rows by review bucket"),
    ]:
        counts = actions[column].value_counts(dropna=False).reset_index()
        counts.columns = ["value", "n_rows"]
        counts.insert(0, "metric", label)
        rows.append(counts)

    mapping_summary = pd.DataFrame(
        [
            {
                "metric": "Validated mapping rows",
                "value": "total",
                "n_rows": len(validated_mapping),
            }
        ]
    )
    rows.append(mapping_summary)
    return pd.concat(rows, ignore_index=True)


def apply_manual_validation(
    main_workbook: Path,
    validation_workbook: Path,
    output_workbook: Path,
) -> dict[str, pd.DataFrame]:
    main_sheets = pd.read_excel(main_workbook, sheet_name=None)
    companies_master = main_sheets["companies_master"]
    mapping_log = main_sheets["mapping_log"]
    tinvest_instruments = main_sheets["tinvest_instruments"]
    tinvest_bonds = main_sheets["tinvest_bonds"]
    moex_instruments = main_sheets["moex_instruments"]
    moex_bonds = main_sheets["moex_bonds"]

    validation_df = load_validation(validation_workbook)
    known_ids = build_known_ids(moex_instruments, moex_bonds, tinvest_instruments)
    actions = build_validated_actions(validation_df, known_ids)

    validated_mapping = build_validated_mapping(
        actions=actions,
        tinvest_instruments=tinvest_instruments,
        moex_instruments=moex_instruments,
        moex_bonds=moex_bonds,
    )

    tinvest_instruments_validated = annotate_tinvest_instruments(tinvest_instruments, actions)
    tinvest_bonds_validated = annotate_tinvest_instruments(tinvest_bonds, actions)
    moex_instruments_validated = add_validated_flag(moex_instruments, validated_mapping, ["secid", "isin"])
    moex_bonds_validated = add_validated_flag(moex_bonds, validated_mapping, ["SECID", "ISIN"])
    validation_summary = build_validation_summary(actions, validated_mapping)

    validation_columns = [
        "company_name",
        "review_bucket",
        "verdict",
        "confidence",
        "validation_action",
        "validated_ids",
        "validated_isins",
        "what_is_correct",
        "what_is_wrong",
        "what_should_be_changed",
        "manual_action_required",
        "notes",
        "source_urls",
    ]

    outputs = {
        "validation_summary": validation_summary,
        "validated_company_actions": actions[[column for column in validation_columns if column in actions.columns]],
        "issuer_mapping_validated": validated_mapping,
        "tinvest_instruments_validated": tinvest_instruments_validated,
        "tinvest_bonds_validated": tinvest_bonds_validated,
        "moex_instruments_validated": moex_instruments_validated,
        "moex_bonds_validated": moex_bonds_validated,
        "manual_validation_checks": validation_df,
        "companies_master": companies_master,
        "mapping_log": mapping_log,
    }

    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_workbook, engine="openpyxl") as writer:
        for sheet_name, frame in outputs.items():
            frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)

    workbook = load_workbook(output_workbook)
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        ws.freeze_panes = "A2"
        for header_cell in ws[1]:
            header_cell.font = Font(bold=True)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_letter in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"]:
            ws.column_dimensions[column_letter].width = 26
    workbook.save(output_workbook)

    return outputs


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Apply manual validation decisions to additional_public_data.xlsx.",
    )
    parser.add_argument("--main-workbook", type=Path, default=MAIN_WORKBOOK_DEFAULT)
    parser.add_argument("--validation-workbook", type=Path, default=VALIDATION_WORKBOOK_DEFAULT)
    parser.add_argument("--output", type=Path, default=OUTPUT_WORKBOOK_DEFAULT)
    args = parser.parse_args()

    outputs = apply_manual_validation(
        main_workbook=args.main_workbook,
        validation_workbook=args.validation_workbook,
        output_workbook=args.output,
    )
    print(f"Validated workbook created: {args.output}")
    for name, frame in outputs.items():
        print(f"{name}: {len(frame)} rows")


if __name__ == "__main__":
    main()
