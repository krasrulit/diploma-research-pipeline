from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from .utils import clean_text, ensure_directory, normalize_identifier, normalize_isin, normalize_text, normalize_ticker, save_dataframe_csv


EXCEL_ROW_LIMIT = 1_000_000
PRACTICAL_WORKBOOK_ROW_LIMIT = 150_000


def bool_value(value: object) -> bool:
    if isinstance(value, bool):
        return value
    text = clean_text(value).lower()
    return text in {"1", "true", "yes", "да"}


def join_unique(values: pd.Series) -> str:
    seen: set[str] = set()
    result: list[str] = []
    for value in values.dropna().astype(str):
        text = clean_text(value)
        if not text or text in seen:
            continue
        seen.add(text)
        result.append(text)
    return " | ".join(result)


def normalize_company_columns(frame: pd.DataFrame) -> pd.DataFrame:
    out = frame.copy()
    if "company_inn" in out.columns:
        out["company_inn"] = out["company_inn"].map(lambda value: normalize_identifier(value, length=10))
    elif "inn" in out.columns:
        out["company_inn"] = out["inn"].map(lambda value: normalize_identifier(value, length=10))

    if "company_ticker" in out.columns:
        out["company_ticker"] = out["company_ticker"].map(normalize_ticker)
    elif "ticker" in out.columns:
        out["company_ticker"] = out["ticker"].map(normalize_ticker)

    if "company_isin" in out.columns:
        out["company_isin"] = out["company_isin"].map(normalize_isin)
    elif "isin" in out.columns:
        out["company_isin"] = out["isin"].map(normalize_isin)
    return out


def company_meta_lookup(companies_master: pd.DataFrame) -> pd.DataFrame:
    if companies_master.empty:
        return pd.DataFrame()
    meta = companies_master.copy()
    meta["inn"] = meta["inn"].map(lambda value: normalize_identifier(value, length=10))
    meta["ticker"] = meta["ticker"].map(normalize_ticker)
    meta["isin"] = meta["isin"].map(normalize_isin)
    meta["company_id"] = meta.apply(
        lambda row: f"{clean_text(row.get('sector', 'unknown'))}:{clean_text(row.get('inn', ''))}"
        if clean_text(row.get("inn", ""))
        else f"{clean_text(row.get('sector', 'unknown'))}:{clean_text(row.get('company_name', ''))}",
        axis=1,
    )
    keep_cols = [
        "company_name",
        "company_name_short",
        "company_name_full",
        "company_name_en",
        "company_name_core",
        "inn",
        "ogrn",
        "ticker",
        "isin",
        "spark_id",
        "industry",
        "sector",
        "sample_flag",
        "sample_membership",
        "selection_result",
        "sample_main_flag",
        "sample_extended_flag",
        "company_id",
    ]
    keep_cols = [column for column in keep_cols if column in meta.columns]
    return meta[keep_cols].drop_duplicates(subset=["inn"], keep="first").reset_index(drop=True)


def coalesce_meta_columns(frame: pd.DataFrame, columns: list[str]) -> pd.DataFrame:
    out = frame.copy()
    for column in columns:
        meta_column = f"{column}_meta"
        if meta_column not in out.columns:
            continue
        if column not in out.columns:
            out[column] = out[meta_column]
            continue

        if pd.api.types.is_bool_dtype(out[meta_column]) or pd.api.types.is_bool_dtype(out[column]):
            out[column] = out[column].fillna(out[meta_column])
        else:
            out[column] = out[column].where(out[column].map(clean_text).ne(""), out[meta_column])
            out[column] = out[column].fillna(out[meta_column])
    drop_cols = [f"{column}_meta" for column in columns if f"{column}_meta" in out.columns]
    if drop_cols:
        out = out.drop(columns=drop_cols)
    return out


def read_csv_if_exists(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    try:
        return pd.read_csv(path, low_memory=False)
    except pd.errors.EmptyDataError:
        return pd.DataFrame()


def security_group_key(company_inn: str, instrument_type: str, ticker: str, isin: str, instrument_name: str) -> str:
    if isin:
        return f"{instrument_type}:isin:{isin}"
    if ticker and company_inn:
        return f"{instrument_type}:inn:{company_inn}:ticker:{ticker}"
    return f"{instrument_type}:inn:{company_inn}:name:{normalize_text(instrument_name)[:120]}"


def ticker_collision_key(company_inn: str, instrument_type: str, ticker: str) -> str:
    if not company_inn or not ticker:
        return ""
    return f"{company_inn}:{instrument_type}:{ticker}"


def first_nonempty(row: pd.Series, columns: list[str]) -> object:
    for column in columns:
        if column not in row.index:
            continue
        value = row.get(column)
        if pd.isna(value):
            continue
        if isinstance(value, str):
            text = clean_text(value)
            if text:
                return text
        else:
            return value
    return np.nan


def prepare_moex_security_master(
    moex_instruments: pd.DataFrame,
    moex_bonds: pd.DataFrame,
    companies_master: pd.DataFrame,
) -> pd.DataFrame:
    if moex_instruments.empty:
        return pd.DataFrame()

    frame = normalize_company_columns(moex_instruments)
    frame["ticker"] = frame.get("secid", pd.Series(index=frame.index, dtype="object")).map(normalize_ticker)
    frame["isin"] = frame.get("isin", pd.Series(index=frame.index, dtype="object")).map(normalize_isin)
    frame["source"] = "moex"
    frame["source_instrument_id"] = frame.get("source_instrument_id", frame.get("secid", "")).map(clean_text)
    frame["source_instrument_id_type"] = "secid"
    frame["instrument_type"] = np.where(
        frame.get("bond_flag", False).map(bool_value),
        "bond",
        np.where(frame.get("share_flag", False).map(bool_value), "share", "other"),
    )
    frame["instrument_name"] = frame.apply(
        lambda row: first_nonempty(row, ["shortname", "name", "emitent_title", "title"]),
        axis=1,
    )
    frame["match_confidence"] = frame.get("confidence", "").map(clean_text)
    frame["match_score"] = pd.to_numeric(frame.get("match_score"), errors="coerce")
    frame["selected_flag"] = frame.get("selected_flag", False).map(bool_value)
    frame["history_candidate_flag"] = frame.get("history_candidate_flag", False).map(bool_value)
    frame["exact_identifier_match"] = frame.get("exact_identifier_match", False).map(bool_value)
    frame["match_reasons"] = frame.get("match_reasons", "").map(clean_text)
    frame["boardid"] = frame.apply(
        lambda row: first_nonempty(row, ["primary_boardid", "primary_board_id", "boardid"]),
        axis=1,
    )
    frame["currency"] = frame.apply(
        lambda row: first_nonempty(row, ["faceunit", "market_faceunit", "currencyid"]),
        axis=1,
    )
    frame["trading_status"] = frame.apply(
        lambda row: first_nonempty(row, ["tradingstatus", "status"]),
        axis=1,
    )
    frame["maturity_date"] = np.nan
    frame["nominal"] = np.nan
    frame["coupon_rate"] = np.nan
    frame["coupon_frequency"] = np.nan

    if not moex_bonds.empty:
        bond_meta = moex_bonds.copy()
        bond_meta["source_instrument_id"] = bond_meta.get("SECID", "").map(clean_text)
        bond_meta = bond_meta.rename(
            columns={
                "SHORTNAME": "bond_shortname",
                "ISIN": "bond_isin",
                "MATDATE": "bond_maturity_date",
                "FACEVALUE": "bond_nominal",
                "COUPONVALUE": "bond_coupon_rate",
                "COUPONFREQUENCY": "bond_coupon_frequency",
                "BOARDID": "bond_boardid",
                "FACEUNIT": "bond_currency",
            }
        )
        merge_cols = [
            "source_instrument_id",
            "bond_shortname",
            "bond_isin",
            "bond_maturity_date",
            "bond_nominal",
            "bond_coupon_rate",
            "bond_coupon_frequency",
            "bond_boardid",
            "bond_currency",
        ]
        merge_cols = [column for column in merge_cols if column in bond_meta.columns]
        frame = frame.merge(
            bond_meta[merge_cols].drop_duplicates(subset=["source_instrument_id"], keep="first"),
            on="source_instrument_id",
            how="left",
        )
        frame["instrument_name"] = frame["bond_shortname"].combine_first(frame["instrument_name"])
        frame["isin"] = frame["bond_isin"].combine_first(frame["isin"]).map(normalize_isin)
        frame["maturity_date"] = frame["bond_maturity_date"]
        frame["nominal"] = pd.to_numeric(frame["bond_nominal"], errors="coerce")
        frame["coupon_rate"] = pd.to_numeric(frame["bond_coupon_rate"], errors="coerce")
        frame["coupon_frequency"] = pd.to_numeric(frame["bond_coupon_frequency"], errors="coerce")
        frame["boardid"] = frame["bond_boardid"].combine_first(frame["boardid"])
        frame["currency"] = frame["bond_currency"].combine_first(frame["currency"])

    company_meta = company_meta_lookup(companies_master).rename(
        columns={
            "company_name": "company_name_shortlist",
            "ticker": "shortlist_ticker",
            "isin": "shortlist_isin",
            "inn": "company_inn",
        }
    )
    frame = frame.merge(company_meta, on="company_inn", how="left", suffixes=("", "_meta"))
    frame = coalesce_meta_columns(
        frame,
        ["sector", "sample_flag", "sample_membership", "sample_main_flag", "sample_extended_flag", "company_id"],
    )
    frame["company_id"] = frame["company_id"].fillna(
        frame.apply(
            lambda row: f"{clean_text(row.get('sector', 'unknown'))}:{clean_text(row.get('company_inn', ''))}"
            if clean_text(row.get("company_inn", ""))
            else "",
            axis=1,
        )
    )
    frame["canonical_security_key"] = frame.apply(
        lambda row: security_group_key(
            company_inn=clean_text(row.get("company_inn", "")),
            instrument_type=clean_text(row.get("instrument_type", "")),
            ticker=normalize_ticker(row.get("ticker", "")),
            isin=normalize_isin(row.get("isin", "")),
            instrument_name=clean_text(row.get("instrument_name", "")),
        ),
        axis=1,
    )
    frame["ticker_collision_key"] = frame.apply(
        lambda row: ticker_collision_key(
            company_inn=clean_text(row.get("company_inn", "")),
            instrument_type=clean_text(row.get("instrument_type", "")),
            ticker=normalize_ticker(row.get("ticker", "")),
        ),
        axis=1,
    )

    keep_cols = [
        "company_id",
        "company_name",
        "company_name_shortlist",
        "company_inn",
        "sector",
        "sample_flag",
        "sample_membership",
        "sample_main_flag",
        "sample_extended_flag",
        "company_ticker",
        "company_isin",
        "source",
        "source_instrument_id",
        "source_instrument_id_type",
        "instrument_type",
        "ticker",
        "isin",
        "instrument_name",
        "canonical_security_key",
        "ticker_collision_key",
        "boardid",
        "currency",
        "trading_status",
        "maturity_date",
        "nominal",
        "coupon_rate",
        "coupon_frequency",
        "match_score",
        "match_confidence",
        "selected_flag",
        "history_candidate_flag",
        "exact_identifier_match",
        "match_reasons",
        "emitent_inn_exact",
        "ticker_exact",
        "isin_exact",
        "secid",
        "shortname",
        "name",
        "query",
        "query_type",
        "group",
        "type",
        "emitent_title",
        "emitent_inn",
    ]
    keep_cols = [column for column in keep_cols if column in frame.columns]
    return frame[keep_cols].drop_duplicates(
        subset=["company_inn", "source", "source_instrument_id"],
        keep="first",
    ).reset_index(drop=True)


def prepare_tinvest_security_master(
    tinvest_instruments: pd.DataFrame,
    companies_master: pd.DataFrame,
) -> pd.DataFrame:
    if tinvest_instruments.empty:
        return pd.DataFrame()

    frame = normalize_company_columns(tinvest_instruments)
    frame["ticker"] = frame.get("ticker", "").map(normalize_ticker)
    frame["isin"] = frame.get("isin", "").map(normalize_isin)
    frame["source"] = "tinvest"
    frame["source_instrument_id"] = frame.get("source_instrument_id", "").map(clean_text)
    frame["source_instrument_id_type"] = frame.get("source_instrument_id_type", "").map(clean_text)
    frame["instrument_type"] = frame.get("instrument_kind_source", "").map(clean_text)
    frame["instrument_name"] = frame.get("name", "").map(clean_text)
    frame["boardid"] = frame.get("class_code", "").map(clean_text)
    frame["currency"] = frame.get("currency", "").map(clean_text)
    frame["trading_status"] = frame.apply(
        lambda row: first_nonempty(
            row,
            [
                "trading_status",
                "api_trade_available_flag",
            ],
        ),
        axis=1,
    )
    frame["maturity_date"] = frame.get("maturity_date")
    frame["nominal"] = pd.to_numeric(frame.get("nominal"), errors="coerce")
    frame["coupon_rate"] = pd.to_numeric(frame.get("initial_nominal"), errors="coerce")
    frame["coupon_frequency"] = pd.to_numeric(frame.get("coupon_quantity_per_year"), errors="coerce")
    frame["match_score"] = pd.to_numeric(frame.get("match_score"), errors="coerce")
    frame["match_confidence"] = frame.get("match_confidence", "").map(clean_text)
    frame["selected_flag"] = frame.get("selected_flag", False).map(bool_value)
    frame["history_candidate_flag"] = frame.get("history_candidate_flag", False).map(bool_value)
    frame["exact_identifier_match"] = frame.get("exact_identifier_match", False).map(bool_value)
    frame["match_reasons"] = frame.get("match_reasons", "").map(clean_text)

    company_meta = company_meta_lookup(companies_master).rename(
        columns={
            "company_name": "company_name_shortlist",
            "ticker": "shortlist_ticker",
            "isin": "shortlist_isin",
            "inn": "company_inn",
        }
    )
    frame = frame.merge(company_meta, on="company_inn", how="left", suffixes=("", "_meta"))
    frame = coalesce_meta_columns(
        frame,
        ["sector", "sample_flag", "sample_membership", "sample_main_flag", "sample_extended_flag", "company_id"],
    )
    frame["company_id"] = frame["company_id"].fillna(
        frame.apply(
            lambda row: f"{clean_text(row.get('sector', 'unknown'))}:{clean_text(row.get('company_inn', ''))}"
            if clean_text(row.get("company_inn", ""))
            else "",
            axis=1,
        )
    )
    frame["canonical_security_key"] = frame.apply(
        lambda row: security_group_key(
            company_inn=clean_text(row.get("company_inn", "")),
            instrument_type=clean_text(row.get("instrument_type", "")),
            ticker=normalize_ticker(row.get("ticker", "")),
            isin=normalize_isin(row.get("isin", "")),
            instrument_name=clean_text(row.get("instrument_name", "")),
        ),
        axis=1,
    )
    frame["ticker_collision_key"] = frame.apply(
        lambda row: ticker_collision_key(
            company_inn=clean_text(row.get("company_inn", "")),
            instrument_type=clean_text(row.get("instrument_type", "")),
            ticker=normalize_ticker(row.get("ticker", "")),
        ),
        axis=1,
    )

    keep_cols = [
        "company_id",
        "company_name",
        "company_name_shortlist",
        "company_inn",
        "sector",
        "sample_flag",
        "sample_membership",
        "sample_main_flag",
        "sample_extended_flag",
        "company_ticker",
        "company_isin",
        "source",
        "source_instrument_id",
        "source_instrument_id_type",
        "instrument_type",
        "ticker",
        "isin",
        "instrument_name",
        "canonical_security_key",
        "ticker_collision_key",
        "boardid",
        "currency",
        "trading_status",
        "maturity_date",
        "nominal",
        "coupon_rate",
        "coupon_frequency",
        "match_score",
        "match_confidence",
        "selected_flag",
        "history_candidate_flag",
        "exact_identifier_match",
        "match_reasons",
        "figi",
        "uid",
        "position_uid",
        "class_code",
        "exchange",
        "instrument_status",
        "api_trade_available_flag",
        "for_iis_flag",
        "for_qual_investor_flag",
        "amortization_flag",
        "floating_coupon_flag",
        "perpetual_flag",
        "subordinated_flag",
    ]
    keep_cols = [column for column in keep_cols if column in frame.columns]
    return frame[keep_cols].drop_duplicates(
        subset=["company_inn", "source", "source_instrument_id"],
        keep="first",
    ).reset_index(drop=True)


def instrument_lookup(master: pd.DataFrame) -> pd.DataFrame:
    if master.empty:
        return pd.DataFrame(columns=["source", "source_instrument_id"])
    lookup = master.copy()
    lookup["match_score"] = pd.to_numeric(lookup["match_score"], errors="coerce")
    lookup = lookup.sort_values(
        ["source", "source_instrument_id", "exact_identifier_match", "selected_flag", "match_score"],
        ascending=[True, True, False, False, False],
        kind="stable",
    )
    return lookup.drop_duplicates(subset=["source", "source_instrument_id"], keep="first")


def prepare_moex_history(
    moex_bond_history: pd.DataFrame,
    moex_share_history: pd.DataFrame,
    master: pd.DataFrame,
) -> pd.DataFrame:
    lookup = instrument_lookup(master)
    frames: list[pd.DataFrame] = []

    if not moex_bond_history.empty:
        bond = moex_bond_history.copy()
        bond["source"] = "moex"
        bond["source_instrument_id"] = bond.get("SECID", "").map(clean_text)
        bond["instrument_type"] = "bond"
        bond["trade_date"] = pd.to_datetime(bond.get("TRADEDATE"), errors="coerce")
        bond["open_price"] = pd.to_numeric(bond.get("OPEN"), errors="coerce")
        bond["high_price"] = pd.to_numeric(bond.get("HIGH"), errors="coerce")
        bond["low_price"] = pd.to_numeric(bond.get("LOW"), errors="coerce")
        bond["close_price"] = pd.to_numeric(bond.get("CLOSE"), errors="coerce")
        bond["volume"] = pd.to_numeric(bond.get("VOLUME"), errors="coerce")
        bond["value"] = pd.to_numeric(bond.get("VALUE"), errors="coerce")
        bond["yield_close"] = pd.to_numeric(bond.get("YIELDCLOSE"), errors="coerce")
        bond["accint"] = pd.to_numeric(bond.get("ACCINT"), errors="coerce")
        frames.append(bond)

    if not moex_share_history.empty:
        share = moex_share_history.copy()
        share["source"] = "moex"
        share["source_instrument_id"] = share.get("SECID", "").map(clean_text)
        share["instrument_type"] = "share"
        share["trade_date"] = pd.to_datetime(share.get("TRADEDATE"), errors="coerce")
        share["open_price"] = pd.to_numeric(share.get("OPEN"), errors="coerce")
        share["high_price"] = pd.to_numeric(share.get("HIGH"), errors="coerce")
        share["low_price"] = pd.to_numeric(share.get("LOW"), errors="coerce")
        share["close_price"] = pd.to_numeric(share.get("CLOSE"), errors="coerce")
        share["volume"] = pd.to_numeric(share.get("VOLUME"), errors="coerce")
        share["value"] = pd.to_numeric(share.get("VALUE"), errors="coerce")
        share["yield_close"] = np.nan
        share["accint"] = np.nan
        frames.append(share)

    if not frames:
        return pd.DataFrame()

    history = pd.concat(frames, ignore_index=True, sort=False)
    history = history.merge(
        lookup[
            [
                column
                for column in [
                    "source",
                    "source_instrument_id",
                    "company_id",
                    "company_name",
                    "company_inn",
                    "sector",
                    "sample_flag",
                    "ticker",
                    "isin",
                    "instrument_name",
                    "canonical_security_key",
                    "ticker_collision_key",
                    "match_score",
                    "match_confidence",
                    "selected_flag",
                    "exact_identifier_match",
                ]
                if column in lookup.columns
            ]
        ],
        on=["source", "source_instrument_id"],
        how="left",
        suffixes=("", "_meta"),
    )
    history = coalesce_meta_columns(
        history,
        [
            "company_id",
            "company_name",
            "company_inn",
            "sector",
            "sample_flag",
            "ticker",
            "isin",
            "instrument_name",
            "canonical_security_key",
            "ticker_collision_key",
            "match_score",
            "match_confidence",
            "selected_flag",
            "exact_identifier_match",
        ],
    )
    keep_cols = [
        "source",
        "source_instrument_id",
        "instrument_type",
        "company_id",
        "company_name",
        "company_inn",
        "sector",
        "sample_flag",
        "ticker",
        "isin",
        "instrument_name",
        "canonical_security_key",
        "ticker_collision_key",
        "trade_date",
        "open_price",
        "high_price",
        "low_price",
        "close_price",
        "volume",
        "value",
        "yield_close",
        "accint",
        "match_score",
        "match_confidence",
        "selected_flag",
        "exact_identifier_match",
    ]
    keep_cols = [column for column in keep_cols if column in history.columns]
    return history[keep_cols].sort_values(["source_instrument_id", "trade_date"], kind="stable").reset_index(drop=True)


def prepare_tinvest_history(
    tinvest_history: pd.DataFrame,
    master: pd.DataFrame,
) -> pd.DataFrame:
    if tinvest_history.empty:
        return pd.DataFrame()

    lookup = instrument_lookup(master)
    history = tinvest_history.copy()
    history["source"] = "tinvest"
    history["source_instrument_id"] = history.get("source_instrument_id", "").map(clean_text)
    history["instrument_type"] = history.get("instrument_kind_source", "").map(clean_text)
    history["trade_date"] = pd.to_datetime(history.get("time"), errors="coerce").dt.tz_localize(None)
    history["open_price"] = pd.to_numeric(history.get("open"), errors="coerce")
    history["high_price"] = pd.to_numeric(history.get("high"), errors="coerce")
    history["low_price"] = pd.to_numeric(history.get("low"), errors="coerce")
    history["close_price"] = pd.to_numeric(history.get("close"), errors="coerce")
    history["volume"] = pd.to_numeric(history.get("volume"), errors="coerce")
    history["value"] = np.nan
    history["yield_close"] = np.nan
    history["accint"] = np.nan

    history = history.merge(
        lookup[
            [
                column
                for column in [
                    "source",
                    "source_instrument_id",
                    "company_id",
                    "company_name",
                    "company_inn",
                    "sector",
                    "sample_flag",
                    "ticker",
                    "isin",
                    "instrument_name",
                    "canonical_security_key",
                    "ticker_collision_key",
                    "match_score",
                    "match_confidence",
                    "selected_flag",
                    "exact_identifier_match",
                ]
                if column in lookup.columns
            ]
        ],
        on=["source", "source_instrument_id"],
        how="left",
        suffixes=("", "_meta"),
    )
    history = coalesce_meta_columns(
        history,
        [
            "company_id",
            "company_name",
            "company_inn",
            "sector",
            "sample_flag",
            "ticker",
            "isin",
            "instrument_name",
            "canonical_security_key",
            "ticker_collision_key",
            "match_score",
            "match_confidence",
            "selected_flag",
            "exact_identifier_match",
        ],
    )
    keep_cols = [
        "source",
        "source_instrument_id",
        "instrument_type",
        "company_id",
        "company_name",
        "company_inn",
        "sector",
        "sample_flag",
        "ticker",
        "isin",
        "instrument_name",
        "canonical_security_key",
        "ticker_collision_key",
        "trade_date",
        "open_price",
        "high_price",
        "low_price",
        "close_price",
        "volume",
        "value",
        "yield_close",
        "accint",
        "match_score",
        "match_confidence",
        "selected_flag",
        "exact_identifier_match",
    ]
    keep_cols = [column for column in keep_cols if column in history.columns]
    return history[keep_cols].sort_values(["source_instrument_id", "trade_date"], kind="stable").reset_index(drop=True)


def build_security_history_coverage(
    security_master_all: pd.DataFrame,
    security_history_all: pd.DataFrame,
) -> pd.DataFrame:
    base = security_master_all.copy()
    if base.empty:
        return pd.DataFrame()

    base["company_inn"] = base["company_inn"].map(lambda value: normalize_identifier(value, length=10))
    base["source"] = base["source"].map(clean_text)
    base["source_instrument_id"] = base["source_instrument_id"].map(clean_text)
    base["canonical_security_key"] = base["canonical_security_key"].map(clean_text)
    base["match_score"] = pd.to_numeric(base.get("match_score"), errors="coerce")
    base = base.drop_duplicates(subset=["company_inn", "source", "source_instrument_id"], keep="first")

    if security_history_all.empty:
        coverage = base.copy()
        coverage["history_start"] = pd.NaT
        coverage["history_end"] = pd.NaT
        coverage["n_rows"] = 0
        coverage["n_trade_dates"] = 0
        coverage["n_close_nonnull"] = 0
        coverage["n_volume_nonnull"] = 0
        coverage["n_value_nonnull"] = 0
        coverage["history_span_days"] = 0
        coverage["history_span_years"] = 0.0
        coverage["close_fill_rate"] = 0.0
        coverage["volume_fill_rate"] = 0.0
        coverage["value_fill_rate"] = 0.0
        coverage["has_recent_data_flag"] = False
        coverage["coverage_score"] = 0.0
        return coverage

    history = security_history_all.copy()
    history["company_inn"] = history["company_inn"].map(lambda value: normalize_identifier(value, length=10))
    history["source"] = history["source"].map(clean_text)
    history["source_instrument_id"] = history["source_instrument_id"].map(clean_text)
    history["canonical_security_key"] = history["canonical_security_key"].map(clean_text)
    history["trade_date"] = pd.to_datetime(history["trade_date"], errors="coerce")
    history = history.loc[history["trade_date"].notna()].copy()
    history["trade_day"] = history["trade_date"].dt.normalize()

    grouped = (
        history.groupby(
            ["source", "source_instrument_id", "company_inn", "canonical_security_key"],
            dropna=False,
        )
        .agg(
            history_start=("trade_day", "min"),
            history_end=("trade_day", "max"),
            n_rows=("trade_day", "size"),
            n_trade_dates=("trade_day", pd.Series.nunique),
            n_close_nonnull=("close_price", lambda values: int(pd.Series(values).notna().sum())),
            n_volume_nonnull=("volume", lambda values: int(pd.Series(values).notna().sum())),
            n_value_nonnull=("value", lambda values: int(pd.Series(values).notna().sum())),
        )
        .reset_index()
    )

    grouped["history_span_days"] = (
        (grouped["history_end"] - grouped["history_start"]).dt.days.fillna(0).astype(int) + 1
    )
    grouped["history_span_years"] = grouped["history_span_days"] / 365.25
    grouped["close_fill_rate"] = np.where(grouped["n_rows"].gt(0), grouped["n_close_nonnull"] / grouped["n_rows"], 0.0)
    grouped["volume_fill_rate"] = np.where(grouped["n_rows"].gt(0), grouped["n_volume_nonnull"] / grouped["n_rows"], 0.0)
    grouped["value_fill_rate"] = np.where(grouped["n_rows"].gt(0), grouped["n_value_nonnull"] / grouped["n_rows"], 0.0)
    cutoff = pd.Timestamp.now("UTC").tz_localize(None).normalize() - pd.Timedelta(days=365)
    grouped["has_recent_data_flag"] = grouped["history_end"].ge(cutoff)

    trade_score = np.clip(grouped["n_trade_dates"] / 750.0, 0.0, 1.0)
    span_score = np.clip(grouped["history_span_days"] / 3650.0, 0.0, 1.0)
    recent_score = grouped["has_recent_data_flag"].astype(float)
    grouped["coverage_score"] = (
        100.0
        * (0.45 * trade_score + 0.25 * span_score + 0.20 * grouped["close_fill_rate"] + 0.10 * recent_score)
    ).round(2)

    coverage = base.merge(
        grouped,
        on=["source", "source_instrument_id", "company_inn", "canonical_security_key"],
        how="left",
    )
    numeric_fill_zero = [
        "n_rows",
        "n_trade_dates",
        "n_close_nonnull",
        "n_volume_nonnull",
        "n_value_nonnull",
        "history_span_days",
        "history_span_years",
        "close_fill_rate",
        "volume_fill_rate",
        "value_fill_rate",
        "coverage_score",
    ]
    for column in numeric_fill_zero:
        coverage[column] = pd.to_numeric(coverage[column], errors="coerce").fillna(0)
    coverage["has_recent_data_flag"] = coverage["has_recent_data_flag"].fillna(False)
    return coverage.sort_values(
        ["company_inn", "instrument_type", "ticker", "source", "coverage_score", "match_score"],
        ascending=[True, True, True, True, False, False],
        kind="stable",
    ).reset_index(drop=True)


def build_security_source_resolution(coverage: pd.DataFrame) -> pd.DataFrame:
    if coverage.empty:
        return pd.DataFrame()

    rows: list[dict[str, object]] = []
    group_cols = ["company_inn", "instrument_type", "canonical_security_key"]

    for _, group in coverage.groupby(group_cols, dropna=False, sort=False):
        work = group.copy()
        work["match_score"] = pd.to_numeric(work["match_score"], errors="coerce").fillna(-1)
        work = work.sort_values(
            ["coverage_score", "n_trade_dates", "exact_identifier_match", "selected_flag", "match_score"],
            ascending=[False, False, False, False, False],
            kind="stable",
        ).reset_index(drop=True)

        best = work.iloc[0]
        second = work.iloc[1] if len(work) > 1 else None
        best_score = float(best.get("coverage_score", 0.0))
        second_score = float(second.get("coverage_score", 0.0)) if second is not None else np.nan
        best_dates = int(best.get("n_trade_dates", 0))
        second_dates = int(second.get("n_trade_dates", 0)) if second is not None else 0
        score_gap = best_score - second_score if second is not None else np.nan

        manual_review = False
        resolution_reason = "single_source"
        if len(work) > 1:
            resolution_reason = "best_coverage_source"
            if (
                second is not None
                and best_dates > 0
                and second_dates > 0
                and abs(score_gap) <= 5
                and (best_dates / max(second_dates, 1)) <= 1.15
            ):
                manual_review = True
                resolution_reason = "coverage_tie_manual_review"
            elif best_dates == 0 and second_dates == 0:
                manual_review = True
                resolution_reason = "no_history_any_source"

        rows.append(
            {
                "company_id": clean_text(best.get("company_id", "")),
                "company_name": clean_text(best.get("company_name_shortlist", best.get("company_name", ""))),
                "company_inn": clean_text(best.get("company_inn", "")),
                "sector": clean_text(best.get("sector", "")),
                "sample_flag": clean_text(best.get("sample_flag", "")),
                "instrument_type": clean_text(best.get("instrument_type", "")),
                "canonical_security_key": clean_text(best.get("canonical_security_key", "")),
                "ticker": clean_text(best.get("ticker", "")),
                "isin": clean_text(best.get("isin", "")),
                "instrument_name": clean_text(best.get("instrument_name", "")),
                "sources_available": join_unique(work["source"]),
                "source_rows_available": len(work),
                "best_source": clean_text(best.get("source", "")),
                "best_source_instrument_id": clean_text(best.get("source_instrument_id", "")),
                "best_coverage_score": best_score,
                "best_n_trade_dates": best_dates,
                "best_history_start": best.get("history_start"),
                "best_history_end": best.get("history_end"),
                "best_match_score": float(best.get("match_score", np.nan)) if pd.notna(best.get("match_score")) else np.nan,
                "alt_source": clean_text(second.get("source", "")) if second is not None else "",
                "alt_source_instrument_id": clean_text(second.get("source_instrument_id", "")) if second is not None else "",
                "alt_coverage_score": second_score,
                "alt_n_trade_dates": second_dates,
                "manual_review_needed_flag": manual_review,
                "resolution_reason": resolution_reason,
                "history_available_flag": best_dates > 0,
            }
        )

    return pd.DataFrame(rows).sort_values(
        ["company_inn", "instrument_type", "ticker", "best_coverage_score"],
        ascending=[True, True, True, False],
        kind="stable",
    ).reset_index(drop=True)


def build_security_manual_review(
    coverage: pd.DataFrame,
    resolution: pd.DataFrame,
) -> pd.DataFrame:
    review_rows: list[dict[str, object]] = []

    if not resolution.empty:
        for _, row in resolution.loc[resolution["manual_review_needed_flag"].fillna(False)].iterrows():
            review_rows.append(
                {
                    "company_name": clean_text(row.get("company_name", "")),
                    "company_inn": clean_text(row.get("company_inn", "")),
                    "sector": clean_text(row.get("sector", "")),
                    "sample_flag": clean_text(row.get("sample_flag", "")),
                    "instrument_type": clean_text(row.get("instrument_type", "")),
                    "ticker": clean_text(row.get("ticker", "")),
                    "isin": clean_text(row.get("isin", "")),
                    "review_scope": "source_resolution",
                    "review_reason": clean_text(row.get("resolution_reason", "")),
                    "candidates_summary": (
                        f"best={clean_text(row.get('best_source', ''))}:{clean_text(row.get('best_source_instrument_id', ''))} "
                        f"({row.get('best_n_trade_dates', 0)} dates); "
                        f"alt={clean_text(row.get('alt_source', ''))}:{clean_text(row.get('alt_source_instrument_id', ''))} "
                        f"({row.get('alt_n_trade_dates', 0)} dates)"
                    ),
                    "suggested_action": "Check both sources manually before choosing the primary history source.",
                }
            )

    if not coverage.empty:
        ticker_collisions = (
            coverage.loc[coverage["ticker_collision_key"].map(clean_text).ne("")]
            .groupby(["ticker_collision_key"], dropna=False)
            .agg(
                company_name=("company_name_shortlist", lambda values: join_unique(pd.Series(values))),
                company_inn=("company_inn", lambda values: join_unique(pd.Series(values))),
                sector=("sector", lambda values: join_unique(pd.Series(values))),
                sample_flag=("sample_flag", lambda values: join_unique(pd.Series(values))),
                instrument_type=("instrument_type", lambda values: join_unique(pd.Series(values))),
                ticker=("ticker", lambda values: join_unique(pd.Series(values))),
                distinct_security_keys=("canonical_security_key", pd.Series.nunique),
                candidate_rows=("source_instrument_id", "size"),
                isins=("isin", lambda values: join_unique(pd.Series(values))),
                sources=("source", lambda values: join_unique(pd.Series(values))),
                instrument_names=("instrument_name", lambda values: join_unique(pd.Series(values))),
                max_trade_dates=("n_trade_dates", "max"),
            )
            .reset_index()
        )
        ticker_collisions = ticker_collisions.loc[
            ticker_collisions["distinct_security_keys"].gt(1) & ticker_collisions["max_trade_dates"].gt(0)
        ].copy()

        for _, row in ticker_collisions.iterrows():
            review_rows.append(
                {
                    "company_name": clean_text(row.get("company_name", "")),
                    "company_inn": clean_text(row.get("company_inn", "")),
                    "sector": clean_text(row.get("sector", "")),
                    "sample_flag": clean_text(row.get("sample_flag", "")),
                    "instrument_type": clean_text(row.get("instrument_type", "")),
                    "ticker": clean_text(row.get("ticker", "")),
                    "isin": clean_text(row.get("isins", "")),
                    "review_scope": "ticker_collision",
                    "review_reason": "multiple_distinct_instruments_same_ticker",
                    "candidates_summary": (
                        f"candidate_rows={int(row.get('candidate_rows', 0))}; "
                        f"distinct_keys={int(row.get('distinct_security_keys', 0))}; "
                        f"sources={clean_text(row.get('sources', ''))}; "
                        f"names={clean_text(row.get('instrument_names', ''))}"
                    ),
                    "suggested_action": "If several real instruments share one ticker, keep both; otherwise mark the wrong mapping for exclusion.",
                }
            )

    if not review_rows:
        return pd.DataFrame()

    review = pd.DataFrame(review_rows).drop_duplicates(
        subset=["company_inn", "instrument_type", "ticker", "review_scope", "review_reason"],
        keep="first",
    )
    return review.sort_values(["company_name", "instrument_type", "ticker"], kind="stable").reset_index(drop=True)


def build_company_security_summary(
    companies_master: pd.DataFrame,
    security_master_all: pd.DataFrame,
    coverage: pd.DataFrame,
    resolution: pd.DataFrame,
) -> pd.DataFrame:
    base = company_meta_lookup(companies_master).rename(
        columns={"inn": "company_inn", "ticker": "company_ticker", "isin": "company_isin"}
    )
    if "company_name" not in base.columns:
        base["company_name"] = ""

    if not security_master_all.empty:
        candidate_counts = (
            security_master_all.groupby(["company_id", "company_name_shortlist", "company_inn", "sector", "sample_flag"], dropna=False)
            .agg(
                n_candidate_rows=("source_instrument_id", "size"),
                n_candidate_shares=("instrument_type", lambda values: int(pd.Series(values).eq("share").sum())),
                n_candidate_bonds=("instrument_type", lambda values: int(pd.Series(values).eq("bond").sum())),
                n_history_candidates=("history_candidate_flag", lambda values: int(pd.Series(values).map(bool_value).sum())),
                n_exact_identifier_matches=("exact_identifier_match", lambda values: int(pd.Series(values).map(bool_value).sum())),
            )
            .reset_index()
            .rename(columns={"company_name_shortlist": "company_name"})
        )
        base = base.merge(candidate_counts, on=["company_id", "company_name", "company_inn", "sector", "sample_flag"], how="left")

    if not resolution.empty:
        resolved = resolution.groupby(["company_id", "company_name", "company_inn", "sector", "sample_flag"], dropna=False).agg(
            n_resolved_groups=("canonical_security_key", "size"),
            n_share_groups=("instrument_type", lambda values: int(pd.Series(values).eq("share").sum())),
            n_bond_groups=("instrument_type", lambda values: int(pd.Series(values).eq("bond").sum())),
            n_groups_with_history=("history_available_flag", lambda values: int(pd.Series(values).fillna(False).sum())),
            n_manual_review_groups=("manual_review_needed_flag", lambda values: int(pd.Series(values).fillna(False).sum())),
            n_best_source_moex=("best_source", lambda values: int(pd.Series(values).eq("moex").sum())),
            n_best_source_tinvest=("best_source", lambda values: int(pd.Series(values).eq("tinvest").sum())),
        ).reset_index()
        base = base.merge(resolved, on=["company_id", "company_name", "company_inn", "sector", "sample_flag"], how="left")

    if not coverage.empty:
        cov = coverage.groupby(["company_id", "company_name_shortlist", "company_inn", "sector", "sample_flag"], dropna=False).agg(
            max_trade_dates=("n_trade_dates", "max"),
            min_history_start=("history_start", "min"),
            max_history_end=("history_end", "max"),
        ).reset_index().rename(columns={"company_name_shortlist": "company_name"})
        base = base.merge(cov, on=["company_id", "company_name", "company_inn", "sector", "sample_flag"], how="left")

    int_columns = [
        "n_candidate_rows",
        "n_candidate_shares",
        "n_candidate_bonds",
        "n_history_candidates",
        "n_exact_identifier_matches",
        "n_resolved_groups",
        "n_share_groups",
        "n_bond_groups",
        "n_groups_with_history",
        "n_manual_review_groups",
        "n_best_source_moex",
        "n_best_source_tinvest",
        "max_trade_dates",
    ]
    for column in int_columns:
        if column in base.columns:
            base[column] = pd.to_numeric(base[column], errors="coerce").fillna(0).astype(int)
    return base.sort_values(["sector", "sample_flag", "company_name"], kind="stable").reset_index(drop=True)


def build_sheet_inventory(outputs: dict[str, pd.DataFrame], workbook_sheets: list[str]) -> pd.DataFrame:
    rows = []
    for sheet_name, frame in outputs.items():
        rows.append(
            {
                "sheet_name": sheet_name,
                "n_rows": len(frame),
                "n_columns": len(frame.columns),
                "written_to_workbook": sheet_name in workbook_sheets,
                "excel_row_limit_exceeded": len(frame) > EXCEL_ROW_LIMIT,
            }
        )
    return pd.DataFrame(rows)


def format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        ws.freeze_panes = "A2"
        for header_cell in ws[1]:
            header_cell.font = Font(bold=True)
            header_cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_letter in ["A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L"]:
            ws.column_dimensions[column_letter].width = 22
    workbook.save(path)


def build_public_securities_outputs(
    companies_master: pd.DataFrame,
    moex_instruments: pd.DataFrame,
    moex_bonds: pd.DataFrame,
    moex_bond_history: pd.DataFrame,
    moex_share_history: pd.DataFrame,
    tinvest_instruments: pd.DataFrame,
    tinvest_bonds: pd.DataFrame,
    tinvest_bond_coupons: pd.DataFrame,
    tinvest_history: pd.DataFrame,
    mapping_log: pd.DataFrame,
    download_log: pd.DataFrame,
) -> dict[str, pd.DataFrame]:
    moex_master = prepare_moex_security_master(moex_instruments, moex_bonds, companies_master)
    tinvest_master = prepare_tinvest_security_master(tinvest_instruments, companies_master)
    security_master_all = pd.concat([moex_master, tinvest_master], ignore_index=True, sort=False)
    if not security_master_all.empty:
        security_master_all = security_master_all.loc[
            security_master_all["instrument_type"].isin(["share", "bond"])
        ].copy()
    if not security_master_all.empty:
        security_master_all = security_master_all.sort_values(
            ["company_inn", "instrument_type", "ticker", "source", "match_score"],
            ascending=[True, True, True, True, False],
            kind="stable",
        ).reset_index(drop=True)

    security_history_moex = prepare_moex_history(moex_bond_history, moex_share_history, security_master_all)
    security_history_tinvest = prepare_tinvest_history(tinvest_history, security_master_all)
    security_history_all = pd.concat(
        [security_history_moex, security_history_tinvest],
        ignore_index=True,
        sort=False,
    ) if (not security_history_moex.empty or not security_history_tinvest.empty) else pd.DataFrame()

    security_history_coverage = build_security_history_coverage(security_master_all, security_history_all)
    security_source_resolution = build_security_source_resolution(security_history_coverage)
    security_manual_review = build_security_manual_review(security_history_coverage, security_source_resolution)
    company_security_summary = build_company_security_summary(
        companies_master,
        security_master_all,
        security_history_coverage,
        security_source_resolution,
    )

    outputs = {
        "companies_master": companies_master,
        "security_master_all": security_master_all,
        "security_history_coverage": security_history_coverage,
        "security_source_resolution": security_source_resolution,
        "security_manual_review": security_manual_review,
        "company_security_summary": company_security_summary,
        "moex_instruments_all": moex_instruments,
        "moex_bonds_all": moex_bonds,
        "tinvest_instruments_all": tinvest_instruments,
        "tinvest_bonds_selected": tinvest_bonds,
        "tinvest_bond_coupons": tinvest_bond_coupons,
        "security_history_moex": security_history_moex,
        "security_history_tinvest": security_history_tinvest,
        "security_history_all": security_history_all,
        "mapping_log": mapping_log,
        "download_log": download_log,
    }
    return outputs


def write_public_securities_workbook(
    outputs: dict[str, pd.DataFrame],
    output_workbook: Path,
    processed_dir: Path,
) -> tuple[dict[str, pd.DataFrame], pd.DataFrame]:
    ensure_directory(output_workbook.parent)
    ensure_directory(processed_dir)

    workbook_outputs: dict[str, pd.DataFrame] = {}
    workbook_sheet_names: list[str] = []

    for sheet_name, frame in outputs.items():
        if len(frame) > PRACTICAL_WORKBOOK_ROW_LIMIT:
            save_dataframe_csv(frame, processed_dir / f"{sheet_name}.csv")
            continue
        if len(frame) > EXCEL_ROW_LIMIT:
            save_dataframe_csv(frame, processed_dir / f"{sheet_name}.csv")
            continue
        workbook_outputs[sheet_name] = frame
        workbook_sheet_names.append(sheet_name)
        save_dataframe_csv(frame, processed_dir / f"{sheet_name}.csv")

    inventory = build_sheet_inventory(outputs, workbook_sheet_names)
    workbook_outputs["sheet_inventory"] = inventory

    with pd.ExcelWriter(output_workbook, engine="openpyxl") as writer:
        for sheet_name, frame in workbook_outputs.items():
            frame.to_excel(writer, sheet_name=sheet_name[:31], index=False)

    format_workbook(output_workbook)
    save_dataframe_csv(inventory, processed_dir / "sheet_inventory.csv")
    return workbook_outputs, inventory
