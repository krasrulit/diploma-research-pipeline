from __future__ import annotations

import argparse
from pathlib import Path
from urllib.parse import quote

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font


MAIN_WORKBOOK_DEFAULT = Path("data_processed/additional_public_data.xlsx")
OUTPUT_WORKBOOK_DEFAULT = Path("data_processed/additional_public_data_manual_review.xlsx")


def clean_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    if text.lower() == "nan":
        return ""
    return text


def first_non_empty(*values: object) -> str:
    for value in values:
        text = clean_text(value)
        if text:
            return text
    return ""


def summarize_frame(
    frame: pd.DataFrame,
    columns: list[str],
    max_rows: int = 8,
) -> str:
    if frame.empty:
        return ""
    lines: list[str] = []
    for _, row in frame.head(max_rows).iterrows():
        parts = []
        for column in columns:
            value = clean_text(row.get(column, ""))
            if value:
                parts.append(f"{column}={value}")
        if parts:
            lines.append("; ".join(parts))
    return "\n".join(lines)


def build_identifier_text(company_row: pd.Series) -> str:
    parts = []
    for key in ["inn", "ticker", "isin", "ogrn", "sample_flag", "sample_membership"]:
        value = clean_text(company_row.get(key, ""))
        if value:
            parts.append(f"{key}={value}")
    return "; ".join(parts)


def build_mapping_text(mapping_row: pd.Series) -> str:
    parts = []
    for key in [
        "status",
        "message",
        "n_candidates",
        "n_selected",
        "best_score",
        "selected_tickers",
        "selected_isins",
    ]:
        value = clean_text(mapping_row.get(key, ""))
        if value:
            parts.append(f"{key}={value}")
    return "; ".join(parts)


def build_moex_search_url(company_row: pd.Series, mapping_row: pd.Series) -> str:
    query = first_non_empty(
        company_row.get("ticker", ""),
        mapping_row.get("selected_isins", ""),
        company_row.get("company_name", ""),
    )
    if not query:
        return ""
    return f"https://iss.moex.com/iss/securities.json?q={quote(query)}&iss.meta=off"


def build_check_steps(company_name: str) -> tuple[str, str, str]:
    return (
        f"Откройте additional_public_data.xlsx, лист companies_master, фильтр company_name = {company_name}",
        f"Откройте additional_public_data.xlsx, лист mapping_log, фильтры source = tinvest и company_name = {company_name}",
        f"Откройте additional_public_data.xlsx, лист tinvest_instruments, фильтр company_name = {company_name}",
    )


def build_review_issue(mapping_row: pd.Series) -> tuple[str, str]:
    n_selected = float(mapping_row.get("n_selected", 0) or 0)
    selected_tickers = clean_text(mapping_row.get("selected_tickers", ""))
    selected_isins = clean_text(mapping_row.get("selected_isins", ""))

    if n_selected > 0 and (selected_tickers or selected_isins):
        return (
            "T-Invest что-то выбрал, но совпадение нужно подтвердить вручную",
            "Проверьте, что найденные бумаги действительно принадлежат нужному юрлицу, а не материнской/дочерней компании с похожим названием.",
        )
    return (
        "T-Invest нашел только слабые кандидаты по названию",
        "Скорее всего это ложные совпадения по части названия. Проверьте, что их нужно оставить unmatched.",
    )


def build_warn_issue() -> tuple[str, str]:
    return (
        "В shortlist есть тикер, но T-Invest не вернул кандидатов",
        "Сверьте тикер из companies_master с MOEX-листами и решите, это отсутствие покрытия T-Invest или проблема исходного тикера.",
    )


def build_gap_issue() -> tuple[str, str]:
    return (
        "В MOEX облигации найдены, но в T-Invest bond sheet их нет",
        "Сверьте ISIN облигаций из moex_bonds с tinvest_instruments и решите, T-Invest реально не покрывает эти выпуски или маппинг можно усилить.",
    )


def priority_for_review(mapping_row: pd.Series) -> str:
    n_selected = float(mapping_row.get("n_selected", 0) or 0)
    best_score = float(mapping_row.get("best_score", 0) or 0)
    if n_selected > 0 or best_score >= 180:
        return "high"
    if best_score >= 120:
        return "medium"
    return "low"


def priority_for_warn(company_row: pd.Series) -> str:
    return "high" if clean_text(company_row.get("ticker", "")) else "medium"


def priority_for_gap(company_name: str, mapping_df: pd.DataFrame) -> str:
    subset = mapping_df[(mapping_df["source"] == "tinvest") & (mapping_df["company_name"] == company_name)]
    if subset.empty:
        return "medium"
    row = subset.iloc[0]
    if clean_text(row.get("status", "")) == "REVIEW":
        return "high"
    return "medium"


def build_manual_review_workbook(
    input_workbook: Path,
    output_workbook: Path,
) -> dict[str, pd.DataFrame]:
    companies_master = pd.read_excel(input_workbook, sheet_name="companies_master")
    mapping_log = pd.read_excel(input_workbook, sheet_name="mapping_log")
    tinvest_instruments = pd.read_excel(input_workbook, sheet_name="tinvest_instruments")
    tinvest_bonds = pd.read_excel(input_workbook, sheet_name="tinvest_bonds")
    moex_instruments = pd.read_excel(input_workbook, sheet_name="moex_instruments")
    moex_bonds = pd.read_excel(input_workbook, sheet_name="moex_bonds")

    company_index = companies_master.set_index(["company_name", "inn"], drop=False)

    review_rows = []
    review_cases = mapping_log.query("source == 'tinvest' and status == 'REVIEW'").copy()
    for _, mapping_row in review_cases.iterrows():
        company_name = clean_text(mapping_row["company_name"])
        inn = clean_text(mapping_row["inn"])
        company_row = company_index.loc[(company_name, inn)] if (company_name, inn) in company_index.index else pd.Series(dtype=object)

        tinvest_subset = tinvest_instruments[tinvest_instruments["company_name"] == company_name]
        moex_subset = moex_instruments[(moex_instruments["company_name"] == company_name) & (moex_instruments["selected_flag"] == True)]
        moex_bond_subset = moex_bonds[moex_bonds["company_name"] == company_name]

        issue_summary, what_to_verify = build_review_issue(mapping_row)
        check_1, check_2, check_3 = build_check_steps(company_name)

        review_rows.append(
            {
                "priority": priority_for_review(mapping_row),
                "review_bucket": "tinvest_review",
                "company_name": company_name,
                "issue_summary": issue_summary,
                "what_is_in_companies_master": build_identifier_text(company_row),
                "what_is_in_mapping_log": build_mapping_text(mapping_row),
                "what_is_in_tinvest": summarize_frame(
                    tinvest_subset,
                    ["ticker", "isin", "name", "instrument_kind_source", "selected_flag", "match_score", "match_confidence"],
                ),
                "what_is_in_moex": summarize_frame(
                    moex_subset,
                    ["secid", "isin", "type", "group", "match_score"],
                ),
                "what_is_in_moex_bonds": summarize_frame(
                    moex_bond_subset,
                    ["SECID", "ISIN", "SHORTNAME"],
                ),
                "what_to_verify_manually": what_to_verify,
                "recommended_next_step": "Подтвердить или отклонить match вручную.",
                "open_main_workbook": str(input_workbook),
                "check_1": check_1,
                "check_2": check_2,
                "check_3": check_3,
                "official_moex_search_url": build_moex_search_url(company_row, mapping_row),
            }
        )

    warn_rows = []
    warn_cases = mapping_log.query("source == 'tinvest' and status == 'WARN'").copy()
    for _, mapping_row in warn_cases.iterrows():
        company_name = clean_text(mapping_row["company_name"])
        inn = clean_text(mapping_row["inn"])
        if (company_name, inn) not in company_index.index:
            continue
        company_row = company_index.loc[(company_name, inn)]
        if not clean_text(company_row.get("ticker", "")):
            continue

        moex_subset = moex_instruments[(moex_instruments["company_name"] == company_name) & (moex_instruments["selected_flag"] == True)]
        moex_bond_subset = moex_bonds[moex_bonds["company_name"] == company_name]
        issue_summary, what_to_verify = build_warn_issue()
        check_1, check_2, check_3 = build_check_steps(company_name)

        warn_rows.append(
            {
                "priority": priority_for_warn(company_row),
                "review_bucket": "tinvest_warn_with_ticker",
                "company_name": company_name,
                "issue_summary": issue_summary,
                "what_is_in_companies_master": build_identifier_text(company_row),
                "what_is_in_mapping_log": build_mapping_text(mapping_row),
                "what_is_in_tinvest": "",
                "what_is_in_moex": summarize_frame(
                    moex_subset,
                    ["secid", "isin", "type", "group", "match_score"],
                ),
                "what_is_in_moex_bonds": summarize_frame(
                    moex_bond_subset,
                    ["SECID", "ISIN", "SHORTNAME"],
                ),
                "what_to_verify_manually": what_to_verify,
                "recommended_next_step": "Понять, это отсутствие покрытия T-Invest или ошибка исходного тикера.",
                "open_main_workbook": str(input_workbook),
                "check_1": check_1,
                "check_2": check_2,
                "check_3": "Откройте additional_public_data.xlsx, лист moex_instruments, фильтр company_name = "
                + company_name,
                "official_moex_search_url": build_moex_search_url(company_row, mapping_row),
            }
        )

    gap_rows = []
    tinvest_bond_companies = set(tinvest_bonds["company_name"].dropna().astype(str))
    for company_name in sorted(set(moex_bonds["company_name"].dropna().astype(str)) - tinvest_bond_companies):
        moex_bond_subset = moex_bonds[moex_bonds["company_name"] == company_name]
        if moex_bond_subset.empty:
            continue
        inn = first_non_empty(moex_bond_subset.iloc[0].get("company_inn", ""))
        company_row = company_index.loc[(company_name, inn)] if (company_name, inn) in company_index.index else pd.Series(dtype=object)
        mapping_subset = mapping_log[(mapping_log["source"] == "tinvest") & (mapping_log["company_name"] == company_name)]
        mapping_row = mapping_subset.iloc[0] if not mapping_subset.empty else pd.Series(dtype=object)
        tinvest_subset = tinvest_instruments[tinvest_instruments["company_name"] == company_name]

        issue_summary, what_to_verify = build_gap_issue()
        check_1, check_2, check_3 = build_check_steps(company_name)

        gap_rows.append(
            {
                "priority": priority_for_gap(company_name, mapping_log),
                "review_bucket": "moex_bonds_but_no_tinvest_bonds",
                "company_name": company_name,
                "issue_summary": issue_summary,
                "what_is_in_companies_master": build_identifier_text(company_row),
                "what_is_in_mapping_log": build_mapping_text(mapping_row),
                "what_is_in_tinvest": summarize_frame(
                    tinvest_subset,
                    ["ticker", "isin", "name", "instrument_kind_source", "selected_flag", "match_score"],
                ),
                "what_is_in_moex": summarize_frame(
                    moex_instruments[
                        (moex_instruments["company_name"] == company_name)
                        & (moex_instruments["selected_flag"] == True)
                    ],
                    ["secid", "isin", "type", "group", "match_score"],
                ),
                "what_is_in_moex_bonds": summarize_frame(
                    moex_bond_subset,
                    ["SECID", "ISIN", "SHORTNAME"],
                ),
                "what_to_verify_manually": what_to_verify,
                "recommended_next_step": "Решить, нужно ли считать это нормальным непокрытием T-Invest или усиливать маппинг.",
                "open_main_workbook": str(input_workbook),
                "check_1": check_1,
                "check_2": check_2,
                "check_3": "Откройте additional_public_data.xlsx, лист moex_bonds, фильтр company_name = "
                + company_name,
                "official_moex_search_url": build_moex_search_url(company_row, mapping_row),
            }
        )

    all_checks = pd.DataFrame(review_rows + warn_rows + gap_rows)
    if not all_checks.empty:
        priority_order = {"high": 0, "medium": 1, "low": 2}
        all_checks["_priority_order"] = all_checks["priority"].map(priority_order).fillna(9)
        all_checks = all_checks.sort_values(
            ["_priority_order", "review_bucket", "company_name"],
            kind="stable",
        ).drop(columns="_priority_order")

    how_to_use = pd.DataFrame(
        [
            {
                "step": 1,
                "instruction": "Откройте этот файл и начните с листа all_checks.",
                "details": "Он уже отсортирован по priority: high -> medium -> low.",
            },
            {
                "step": 2,
                "instruction": "В колонке issue_summary видно, почему компания попала в ручную проверку.",
                "details": "В колонках what_is_in_* уже собраны значения из разных листов основного workbook.",
            },
            {
                "step": 3,
                "instruction": "Колонки check_1 / check_2 / check_3 говорят, какой лист открыть в основном workbook.",
                "details": "Обычно достаточно открыть additional_public_data.xlsx и поставить фильтр по company_name.",
            },
            {
                "step": 4,
                "instruction": "Колонка official_moex_search_url помогает быстро проверить бумагу во внешнем официальном источнике.",
                "details": "Ссылка ведет на официальный MOEX ISS search endpoint.",
            },
            {
                "step": 5,
                "instruction": "Основной workbook для сверки",
                "details": str(input_workbook),
            },
        ]
    )

    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_workbook, engine="openpyxl") as writer:
        how_to_use.to_excel(writer, sheet_name="how_to_use", index=False)
        all_checks.to_excel(writer, sheet_name="all_checks", index=False)
        pd.DataFrame(review_rows).to_excel(writer, sheet_name="review_cases", index=False)
        pd.DataFrame(warn_rows).to_excel(writer, sheet_name="warn_with_ticker", index=False)
        pd.DataFrame(gap_rows).to_excel(writer, sheet_name="moex_tinvest_gap", index=False)

    workbook = load_workbook(output_workbook)
    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        ws.freeze_panes = "A2"
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for header_cell in ws[1]:
            header_cell.font = Font(bold=True)
        if ws.max_column >= 1:
            ws.column_dimensions["A"].width = 18
        for column_letter in ["B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P"]:
            if column_letter in ws.column_dimensions:
                ws.column_dimensions[column_letter].width = 28

        header_map = {cell.value: cell.column for cell in ws[1] if cell.value}
        if "open_main_workbook" in header_map:
            col_idx = header_map["open_main_workbook"]
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if clean_text(cell.value):
                    cell.hyperlink = f"file://{cell.value}"
                    cell.style = "Hyperlink"
        if "official_moex_search_url" in header_map:
            col_idx = header_map["official_moex_search_url"]
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if clean_text(cell.value):
                    cell.hyperlink = clean_text(cell.value)
                    cell.style = "Hyperlink"

    workbook.save(output_workbook)

    return {
        "all_checks": all_checks,
        "review_cases": pd.DataFrame(review_rows),
        "warn_with_ticker": pd.DataFrame(warn_rows),
        "moex_tinvest_gap": pd.DataFrame(gap_rows),
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build manual review workbook for additional_public_data.xlsx",
    )
    parser.add_argument(
        "--input",
        type=Path,
        default=MAIN_WORKBOOK_DEFAULT,
        help="Path to additional_public_data.xlsx",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=OUTPUT_WORKBOOK_DEFAULT,
        help="Path to the output manual-review workbook",
    )
    args = parser.parse_args()

    build_manual_review_workbook(
        input_workbook=args.input,
        output_workbook=args.output,
    )
    print(f"Manual review workbook created: {args.output}")


if __name__ == "__main__":
    main()
